import importlib.util
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

RAIZ = Path(__file__).resolve().parents[1] / "hermeshq"
HAY_OFICINA = bool(shutil.which("soffice") or shutil.which("libreoffice"))


def _cargar(nombre, ruta):
    spec = importlib.util.spec_from_file_location(nombre, ruta)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


def _a_formato(origen: Path, formato: str, carpeta: Path) -> Path:
    oficina = shutil.which("soffice") or shutil.which("libreoffice")
    subprocess.run(
        [oficina, "--headless", "--convert-to", formato, "--outdir", str(carpeta), str(origen)],
        capture_output=True,
        timeout=120,
        env=dict(os.environ, HOME=str(carpeta)),
        check=False,
    )
    return carpeta / (origen.stem + "." + formato)


@unittest.skipUnless(HAY_OFICINA, "LibreOffice no está instalado")
class TestFormatosAntiguos(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.carpeta = Path(self.tmp.name)
        import openpyxl

        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.title = "Licencias"
        hoja.append(["Sucursal", "Estado", "Costo"])
        hoja.append(["Arica", "Asignada", 10])
        hoja.append(["Angamos", "Disponible", 12])
        self.xlsx = self.carpeta / "datos.xlsx"
        libro.save(self.xlsx)

    def test_sharepoint_reads_an_old_xls_by_converting_it(self):
        plugin = _cargar(
            "sharepoint_legacy_under_test", RAIZ / "integration_packages" / "sharepoint" / "plugin" / "__init__.py"
        )
        xls = _a_formato(self.xlsx, "xls", self.carpeta)
        self.assertTrue(xls.is_file())
        texto, error, extra = plugin._extract_text("Licencias.xls", xls.read_bytes(), 20000)
        self.assertEqual(error, "")
        self.assertIn("Angamos", texto)
        self.assertEqual(extra["convertido_desde"], ".xls")

    def test_sharepoint_reads_an_old_doc(self):
        plugin = _cargar(
            "sharepoint_legacy_doc_under_test", RAIZ / "integration_packages" / "sharepoint" / "plugin" / "__init__.py"
        )
        import docx

        documento = docx.Document()
        documento.add_paragraph("Procedimiento de transferencia de asfalto")
        origen = self.carpeta / "proc.docx"
        documento.save(origen)
        antiguo = _a_formato(origen, "doc", self.carpeta)
        texto, error, _ = plugin._extract_text("proc.doc", antiguo.read_bytes(), 20000)
        self.assertEqual(error, "")
        self.assertIn("transferencia de asfalto", texto)

    def test_without_a_converter_the_message_says_so(self):
        plugin = _cargar(
            "sharepoint_legacy_bad_under_test", RAIZ / "integration_packages" / "sharepoint" / "plugin" / "__init__.py"
        )
        with patch.object(plugin.shutil, "which", return_value=None):
            texto, error, _ = plugin._extract_text("roto.xls", b"x", 20000)
        self.assertIsNone(texto)
        self.assertIn("No hay conversor instalado", error)

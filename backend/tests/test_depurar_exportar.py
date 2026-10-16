import importlib.util
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

RAIZ = Path(__file__).resolve().parents[1] / "hermeshq"
PLUGIN = RAIZ / "plugin_templates" / "hermeshq_documentos"


def _cargar(nombre, ruta):
    spec = importlib.util.spec_from_file_location(nombre, ruta)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


plugin = _cargar("documentos_exportar_under_test", PLUGIN / "__init__.py")
exportar = _cargar("exportar_under_test", PLUGIN / "exportar.py")

GOLPE = "Trabajador se golpea la mano derecha con una herramienta al mover una tuberia en la zona norte del patio"
CAIDA = "Operador resbala en la escalera del muelle y cae sobre su costado sin lesion visible durante el turno"
CASOS = [
    {"N°": 3846, "Sucursal": "Arica", "Fecha": "2026-02-03", "Tipo": "STP", "Desc": GOLPE},
    {"N°": 3887, "Sucursal": "Arica", "Fecha": "2026-02-03", "Tipo": "STP", "Desc": GOLPE.upper()},
    {"N°": 3894, "Sucursal": "Arica", "Fecha": "2026-03-05", "Tipo": "STP", "Desc": CAIDA},
    {"N°": 3895, "Sucursal": "Arica", "Fecha": "2026-03-06", "Tipo": "STP", "Desc": CAIDA},
    {"N°": 3900, "Sucursal": "Arica", "Fecha": "2026-03-20", "Tipo": "Cuasi", "Desc": CAIDA},
    {
        "N°": 3901,
        "Sucursal": "Arica",
        "Fecha": "2026-02-03",
        "Tipo": "Cuasi",
        "Desc": "Trabajador se golpea la mano con una herramienta pero en otro sector del taller por un descuido",
    },
    {"N°": 3902, "Sucursal": "Angamos", "Fecha": "2026-02-03", "Tipo": "STP", "Desc": GOLPE},
    {"N°": 3903, "Sucursal": "Arica", "Fecha": "2026-04-01", "Tipo": "Cuasi", "Desc": "Caída"},
    {"N°": 3904, "Sucursal": "Arica", "Fecha": "2026-04-01", "Tipo": "Cuasi", "Desc": "Caída"},
]


class TestDepurarRepetidos(unittest.TestCase):
    def test_same_day_and_consecutive_numbers_are_removed_keeping_the_oldest(self):
        unicas, descartadas = exportar.depurar_repetidos(CASOS, "N°", "Fecha", "Desc", "Sucursal")
        self.assertEqual({d["descartado"]: d["conservado"] for d in descartadas}, {3887: 3846, 3895: 3894})
        self.assertEqual([c["N°"] for c in unicas], [3846, 3894, 3900, 3901, 3902, 3903, 3904])

    def test_similar_but_different_cases_are_kept(self):
        unicas, _ = exportar.depurar_repetidos(CASOS, "N°", "Fecha", "Desc", "Sucursal")
        numeros = {c["N°"] for c in unicas}
        self.assertIn(3901, numeros)
        self.assertIn(3900, numeros)
        self.assertIn(3902, numeros)
        self.assertIn(3904, numeros)

    def test_rows_without_a_date_or_text_are_never_removed(self):
        filas = [{"N°": 1, "Fecha": "", "Desc": GOLPE}, {"N°": 2, "Fecha": "", "Desc": GOLPE}]
        unicas, descartadas = exportar.depurar_repetidos(filas, "N°", "Fecha", "Desc")
        self.assertEqual((len(unicas), descartadas), (2, []))


class TestExportarExcel(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.raiz = Path(self.tmp.name)
        carpeta = self.raiz / "work" / "hse"
        carpeta.mkdir(parents=True)
        tabla = {
            "nombre": "Inc. Seguridad",
            "filas": len(CASOS),
            "columnas": ["N°", "Sucursal", "Fecha", "Tipo", "Desc"],
            "registros": CASOS,
        }
        (carpeta / "datos.json").write_text(
            json.dumps({"archivos": [{"archivo": "t.xlsx", "hojas": [tabla]}]}), encoding="utf-8"
        )
        entorno = patch.dict(os.environ, {"HERMES_SESSION_CWD": self.tmp.name})
        entorno.start()
        self.addCleanup(entorno.stop)

    def _exportar(self, **args):
        return json.loads(plugin._exportar_datos({"tema": "hse", "hoja": "seguridad", **args}))

    def test_exports_year_month_and_category_with_summary_and_dedup_sheet(self):
        import openpyxl

        r = self._exportar(
            filtros=[{"columna": "Sucursal", "op": "=", "valor": "Arica"}],
            columnas=["N°", "Fecha", "Año", "Mes", "Categoría", "Tipo"],
            columnas_calculadas=[
                {"nombre": "Año", "tipo": "anio", "de": "Fecha"},
                {"nombre": "Mes", "tipo": "mes_nombre", "de": "Fecha"},
                {
                    "nombre": "Categoría",
                    "tipo": "clasificar",
                    "de": "Tipo",
                    "reglas": {"Accidente": ["STP", "CTP"]},
                    "otro": "Incidente",
                },
            ],
            quitar_repetidos={"id": "N°", "fecha": "Fecha", "descripcion": "Desc", "sucursal": "Sucursal"},
            resumen_por=["Categoría", "Mes"],
            ordenar="N°",
            nombre="Casos Arica 2026",
        )
        self.assertTrue(r["success"], r)
        self.assertEqual(r["filas"], 6)
        self.assertEqual(r["repetidos_quitados"], 2)
        self.assertEqual(r["hojas"], ["Inc. Seguridad", "Resumen", "Depuración"])
        self.assertEqual(r["adjuntar_en_la_respuesta"], [f"MEDIA:{r['path']}"])
        libro = openpyxl.load_workbook(self.raiz / r["path"])
        filas = list(libro["Inc. Seguridad"].iter_rows(values_only=True))
        self.assertEqual(filas[0], ("N°", "Fecha", "Año", "Mes", "Categoría", "Tipo"))
        self.assertEqual(filas[1][:5], (3846, "2026-02-03", 2026, "Febrero", "Accidente"))
        self.assertEqual({f[4] for f in filas[1:]}, {"Accidente", "Incidente"})
        resumen = list(libro["Resumen"].iter_rows(values_only=True))
        self.assertIn(("Categoría", "Accidente", 2), resumen)
        self.assertIn(("Mes", "Febrero", 2), resumen)
        depuracion = list(libro["Depuración"].iter_rows(values_only=True))
        self.assertEqual({(f[0], f[1]) for f in depuracion[1:]}, {(3846, 3887), (3894, 3895)})

    def test_errors_are_actionable(self):
        self.assertFalse(self._exportar(hoja="")["success"])
        self.assertIn("no encuentro la columna", self._exportar(columnas=["Inexistente"])["error"])
        self.assertIn(
            "tipo de columna",
            self._exportar(columnas_calculadas=[{"nombre": "X", "tipo": "raro", "de": "Fecha"}])["error"],
        )
        self.assertIn("quitar_repetidos necesita", self._exportar(quitar_repetidos={"id": "N°"})["error"])
        self.assertIn("hq_datos_extraer", json.loads(plugin._exportar_datos({"tema": "otro", "hoja": "x"}))["error"])


if __name__ == "__main__":
    unittest.main()

import importlib.util
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

PLUGIN = Path(__file__).resolve().parents[1] / "hermeshq" / "plugin_templates" / "hermeshq_pptx" / "__init__.py"
_spec = importlib.util.spec_from_file_location("hermeshq_pptx_template_under_test", PLUGIN)
plugin = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(plugin)


class PptxTemplateTests(unittest.TestCase):
    def setUp(self):
        from pptx import Presentation

        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        entorno = patch.dict(os.environ, {"HERMES_SESSION_CWD": self.tmp.name})
        entorno.start()
        self.addCleanup(entorno.stop)
        base = Presentation()
        portada = base.slides.add_slide(base.slide_layouts[0])
        portada.shapes.title.text = "LÁMINA DE EJEMPLO DE LA PLANTILLA"
        base.save(str(self.dir / "plantilla.pptx"))
        from PIL import Image

        Image.new("RGB", (400, 200), (30, 90, 160)).save(self.dir / "foto.png")

    def _llamar(self, **args):
        return json.loads(plugin._handle_from_template({"template": "plantilla.pptx", **args}))

    def _abrir(self, salida):
        from pptx import Presentation

        return Presentation(salida["path"])

    def test_inspecting_lists_layouts_placeholders_and_example_slides(self):
        r = self._llamar(inspect=True)
        self.assertTrue(r["success"], r)
        nombres = [d["nombre"] for d in r["disenos"]]
        self.assertIn("Title Slide", nombres)
        self.assertIn("Title and Content", nombres)
        contenido = next(d for d in r["disenos"] if d["nombre"] == "Title and Content")
        self.assertIn("TITLE", {m["tipo"] for m in contenido["marcadores"]})
        self.assertEqual(r["laminas_existentes"][0]["titulo"], "LÁMINA DE EJEMPLO DE LA PLANTILLA")
        self.assertEqual(self._llamar()["disenos"], r["disenos"])

    def test_slides_use_the_chosen_layouts_and_fill_their_placeholders(self):
        r = self._llamar(
            slides=[
                {"layout": "Title Slide", "title": "Resultados 2026", "subtitle": "Comité"},
                {
                    "layout": "Title and Content",
                    "title": "Hallazgos",
                    "bullets": ["Uno", {"text": "Sub", "level": 1}],
                    "notes": "Decir esto",
                },
            ],
            filename="demo",
        )
        self.assertTrue(r["success"], r)
        self.assertEqual(r["slide_count"], 2)
        self.assertEqual(r["avisos"], [])
        prs = self._abrir(r)
        self.assertEqual(len(prs.slides), 2)
        self.assertEqual(prs.slides[0].shapes.title.text, "Resultados 2026")
        cuerpo = [s for s in prs.slides[1].placeholders if s.placeholder_format.type.name == "OBJECT"][0]
        self.assertEqual([p.text for p in cuerpo.text_frame.paragraphs], ["Uno", "Sub"])
        self.assertEqual(cuerpo.text_frame.paragraphs[1].level, 1)
        self.assertEqual(prs.slides[1].notes_slide.notes_text_frame.text, "Decir esto")
        textos = " ".join(s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame)
        self.assertNotIn("EJEMPLO DE LA PLANTILLA", textos)

    def test_empty_placeholders_are_removed(self):
        r = self._llamar(slides=[{"layout": "Title Slide", "title": "Solo título"}])
        prs = self._abrir(r)
        self.assertEqual([s.placeholder_format.type.name for s in prs.slides[0].placeholders], ["CENTER_TITLE"])

    def test_images_and_tables_are_placed_inside_the_slide(self):
        r = self._llamar(
            slides=[
                {"title": "Con imagen", "image_path": "foto.png"},
                {"title": "Con tabla", "table": {"headers": ["Mes", "Casos"], "rows": [["Ene", 3], ["Feb", 5]]}},
                {"title": "Texto e imagen", "bullets": ["a", "b"], "image_path": "foto.png"},
            ]
        )
        self.assertTrue(r["success"], r)
        prs = self._abrir(r)
        for lamina in prs.slides:
            for forma in lamina.shapes:
                self.assertGreaterEqual(forma.left, 0)
                self.assertLessEqual(forma.left + forma.width, prs.slide_width)
        for numero in (0, 2):
            self.assertTrue(any(hasattr(f, "image") for f in prs.slides[numero].shapes), numero)
        tabla = next(f for f in prs.slides[1].shapes if f.has_table)
        self.assertEqual(tabla.table.cell(0, 0).text, "Mes")
        self.assertEqual(tabla.table.cell(2, 1).text, "5")

    def test_keeping_the_template_slides_appends_after_them(self):
        r = self._llamar(slides=[{"title": "Nueva"}], keep_template_slides=True)
        self.assertEqual(r["slide_count"], 2)
        self.assertEqual(len(self._abrir(r).slides), 2)

    def test_unknown_layouts_and_placeholders_are_explained(self):
        r = self._llamar(slides=[{"layout": "No existe", "title": "x"}])
        self.assertFalse(r["success"])
        self.assertIn("Lámina 1", r["error"])
        self.assertIn("Disponibles", r["error"])
        r = self._llamar(slides=[{"layout": 99, "title": "x"}])
        self.assertIn("no existe", r["error"])
        r = self._llamar(slides=[{"layout": "Title Slide", "title": "x", "fields": {"Marcador inventado": "y"}}])
        self.assertTrue(r["success"])
        self.assertTrue(any("Marcador inventado" in a for a in r["avisos"]))

    def test_bad_input_is_reported(self):
        self.assertFalse(json.loads(plugin._handle_from_template({}))["success"])
        r = json.loads(plugin._handle_from_template({"template": "no-esta.pptx", "inspect": True}))
        self.assertIn("No encuentro la plantilla", r["error"])
        (self.dir / "x.potx").write_bytes(b"x")
        r = json.loads(plugin._handle_from_template({"template": "x.potx", "inspect": True}))
        self.assertIn(".pptx", r["error"])
        self.assertIn("imagen", self._llamar(slides=[{"title": "a", "image_path": "falta.png"}])["error"])

    def test_the_tool_is_registered_and_listed_in_the_manifest(self):
        class Ctx:
            def __init__(self):
                self.nombres = []

            def register_tool(self, **kwargs):
                self.nombres.append(kwargs["name"])

        ctx = Ctx()
        plugin.register(ctx)
        self.assertEqual(sorted(ctx.nombres), ["hq_pptx_from_template", "hq_pptx_generate"])
        manifiesto = PLUGIN.with_name("plugin.yaml").read_text(encoding="utf-8")
        for nombre in ctx.nombres:
            self.assertIn(nombre, manifiesto)


if __name__ == "__main__":
    unittest.main()

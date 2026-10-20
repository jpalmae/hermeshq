import importlib.util
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

_PLUGIN = (
    Path(__file__).resolve().parents[1] / "hermeshq" / "integration_packages" / "sharepoint" / "plugin" / "__init__.py"
)
_spec = importlib.util.spec_from_file_location("sharepoint_plugin", _PLUGIN)
sp = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(sp)


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Consentimiento informado")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Nombre"
    table.rows[0].cells[1].text = "Maria"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["a", "b"])
    ws.append([1, 2])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Titulo uno"
    slide.placeholders[1].text = "Cuerpo"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf_bytes() -> bytes | None:
    try:
        import pypdf  # noqa: F401
        from weasyprint import HTML
    except ImportError:
        return None
    return HTML(string="<p>Texto del consentimiento</p>").write_pdf()


class TestExtraction(unittest.TestCase):
    def test_text(self):
        text, err, _ = sp._extract_text("a.txt", "hola ñandú".encode(), 1000)
        self.assertEqual(text, "hola ñandú")
        self.assertEqual(err, "")

    def test_docx(self):
        text, err, _ = sp._extract_text("a.docx", _docx_bytes(), 1000)
        self.assertIn("Consentimiento informado", text)
        self.assertIn("Nombre | Maria", text)

    def test_xlsx(self):
        text, _, _ = sp._extract_text("a.xlsx", _xlsx_bytes(), 1000)
        self.assertIn("Hoja: Datos", text)
        self.assertIn("1 | 2", text)

    def test_pptx(self):
        text, _, _ = sp._extract_text("a.pptx", _pptx_bytes(), 1000)
        self.assertIn("Diapositiva 1", text)
        self.assertIn("Titulo uno", text)

    def test_pdf(self):
        data = _pdf_bytes()
        if data is None:
            self.skipTest("pypdf/weasyprint not installed")
        text, err, extra = sp._extract_text("a.pdf", data, 1000)
        self.assertIn("Texto del consentimiento", text)
        self.assertEqual(extra["pages"], 1)

    def test_corrupt_pdf_is_friendly(self):
        text, err, _ = sp._extract_text("a.pdf", b"no es un pdf", 1000)
        self.assertIsNone(text)
        self.assertTrue(err)

    def test_unsupported(self):
        text, err, _ = sp._extract_text("a.zip", b"PK", 1000)
        self.assertIsNone(text)
        self.assertIn("no soportado", err)


class TestReadFileTool(unittest.TestCase):
    def _run(self, args, meta, data=b"contenido", user="u1", token="tok"):
        patches = [
            patch.object(sp, "_task_user_id", return_value=user),
            patch.object(sp, "_get_m365_token", return_value=(token, None, "" if token else "sin cuenta")),
            patch.object(sp, "_graph", return_value=meta),
            patch.object(sp, "_download_bytes", return_value=(data, "") if data is not None else (None, "fallo")),
        ]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)
        return json.loads(sp._read_file_tool(args))

    def test_requires_user(self):
        out = self._run({"file_path": "a.txt"}, {}, user=None)
        self.assertFalse(out["success"])
        self.assertIn("usuario", out["error"])

    def test_requires_path_or_id(self):
        out = self._run({}, {})
        self.assertFalse(out["success"])

    def test_graph_error_propagates(self):
        out = self._run({"file_path": "x.txt"}, {"error": "Archivo no encontrado (404)."})
        self.assertTrue(out["error"].startswith("Archivo no encontrado (404)."))
        self.assertIn("sharepoint_search", out["error"])

    def test_folder_rejected(self):
        out = self._run({"file_path": "Docs"}, {"name": "Docs", "folder": {}})
        self.assertIn("carpeta", out["error"])

    def test_too_big(self):
        out = self._run({"file_path": "a.txt"}, {"name": "a.txt", "size": 11 * 1024 * 1024})
        self.assertFalse(out["success"])
        self.assertIn("límite", out["error"])

    def test_success_and_truncation(self):
        body = ("x" * 3000).encode()
        out = self._run({"file_path": "a.txt", "max_chars": 500}, {"name": "a.txt", "size": 3000}, data=body)
        self.assertTrue(out["success"])
        self.assertTrue(out["truncated"])
        self.assertEqual(len(out["content"]), 500)

    def test_download_failure(self):
        out = self._run({"file_path": "a.txt"}, {"name": "a.txt", "size": 5}, data=None)
        self.assertFalse(out["success"])
        self.assertEqual(out["error"], "fallo")

    def test_item_id_path(self):
        with (
            patch.object(sp, "_graph", return_value={"name": "a.txt", "size": 1}) as graph,
            patch.object(sp, "_task_user_id", return_value="u1"),
            patch.object(sp, "_get_m365_token", return_value=("tok", None, "")),
            patch.object(sp, "_download_bytes", return_value=(b"hi", "")),
        ):
            out = json.loads(sp._read_file_tool({"item_id": "ABC", "drive_id": "D1"}))
        self.assertTrue(out["success"])
        self.assertEqual(graph.call_args.args[1], "/drives/D1/items/ABC")

    def test_no_auth_header_on_download_url(self):
        self._run(
            {"file_path": "a.txt"},
            {"name": "a.txt", "size": 1, "@microsoft.graph.downloadUrl": "https://x/y"},
            data=b"x",
        )
        self.assertEqual(sp._download_bytes.call_args.args, ("https://x/y",))


class TestImagesAndDownload(unittest.TestCase):
    def setUp(self):
        self._cwd = os.getcwd()
        self._tmp = tempfile.TemporaryDirectory()
        os.chdir(self._tmp.name)
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        os.chdir(self._cwd)
        self._tmp.cleanup()

    def _run(self, tool, args, meta, data=b"\x89PNG-bytes"):
        with (
            patch.object(sp, "_task_user_id", return_value="u1"),
            patch.object(sp, "_get_m365_token", return_value=("tok", None, "")),
            patch.object(sp, "_graph", return_value=meta),
            patch.object(sp, "_download_bytes", return_value=(data, "")),
        ):
            return json.loads(tool(args))

    def test_read_image_saves_to_workspace(self):
        out = self._run(sp._read_file_tool, {"file_path": "Docs/foto.png"}, {"name": "foto.png", "size": 10})
        self.assertTrue(out["success"])
        self.assertTrue(out["local_path"].startswith("sharepoint_downloads/"))
        self.assertTrue(out["local_path"].endswith("_foto.png"))
        self.assertEqual(out["mime_type"], "image/png")
        self.assertIn("vision_analyze", out["hint"])
        self.assertEqual((Path(self._tmp.name) / out["local_path"]).read_bytes(), b"\x89PNG-bytes")
        self.assertNotIn("content", out)

    def test_download_any_file(self):
        out = self._run(sp._download_file_tool, {"file_path": "a.xlsx"}, {"name": "a.xlsx", "size": 10}, data=b"PK")
        self.assertTrue(out["success"])
        self.assertTrue((Path(self._tmp.name) / out["local_path"]).exists())
        self.assertNotIn("hint", out)

    def test_filename_cannot_escape_download_dir(self):
        out = self._run(sp._download_file_tool, {"file_path": "x"}, {"name": "../../etc/passwd.png", "size": 10})
        resolved = (Path(self._tmp.name) / out["local_path"]).resolve()
        self.assertTrue(str(resolved).startswith(str(Path(self._tmp.name).resolve() / "sharepoint_downloads")))
        self.assertTrue(out["local_path"].endswith("_passwd.png"))

    def test_download_respects_size_limit(self):
        out = self._run(sp._download_file_tool, {"file_path": "x"}, {"name": "big.png", "size": 11 * 1024 * 1024})
        self.assertFalse(out["success"])
        self.assertFalse((Path(self._tmp.name) / "sharepoint_downloads").exists())

    def test_download_requires_path(self):
        out = self._run(sp._download_file_tool, {}, {})
        self.assertFalse(out["success"])


if __name__ == "__main__":
    unittest.main()

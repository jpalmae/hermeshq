import difflib
import re
import unicodedata
from datetime import date

MESES = [
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
]
UMBRAL_MISMO_DIA = 0.92
UMBRAL_CONSECUTIVO = 0.97
MIN_PALABRAS = 4
MAX_FILAS = 20000


def sin_acentos(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor if valor is not None else "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto.lower()).strip()


def fecha_de(valor):
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(valor or ""))
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _entero(valor):
    try:
        return int(float(str(valor).strip().replace(",", ".")))
    except (ValueError, TypeError):
        return None


def _parecido(a, b, umbral):
    if not a or not b or len(a.split()) < MIN_PALABRAS or len(b.split()) < MIN_PALABRAS:
        return False
    comparador = difflib.SequenceMatcher(None, a, b, autojunk=False)
    return (
        comparador.real_quick_ratio() >= umbral and comparador.quick_ratio() >= umbral and comparador.ratio() >= umbral
    )


def depurar_repetidos(filas, col_id=None, col_fecha=None, col_descripcion=None, col_grupo=None):
    conservadas, por_dia, descartadas, retiradas = [], {}, [], set()
    claves = []
    for indice, fila in enumerate(filas):
        numero = _entero(fila.get(col_id)) if col_id else None
        claves.append((numero if numero is not None else float("inf"), indice))
    for _, indice in sorted(claves):
        fila = filas[indice]
        dia = fecha_de(fila.get(col_fecha))
        texto = sin_acentos(fila.get(col_descripcion)) if col_descripcion else ""
        grupo = sin_acentos(fila.get(col_grupo)) if col_grupo else ""
        numero = _entero(fila.get(col_id)) if col_id else None
        repetida = None
        if dia is not None and texto:
            for delta in (0, -1, 1):
                for previa in por_dia.get((grupo, dia.toordinal() + delta), []):
                    otra = filas[previa]
                    texto_previo = sin_acentos(otra.get(col_descripcion))
                    if delta == 0:
                        if _parecido(texto, texto_previo, UMBRAL_MISMO_DIA):
                            repetida = (previa, "mismo día, sucursal y descripción prácticamente idéntica")
                    else:
                        numero_previo = _entero(otra.get(col_id)) if col_id else None
                        if (
                            numero is not None
                            and numero_previo is not None
                            and abs(numero - numero_previo) <= 2
                            and _parecido(texto, texto_previo, UMBRAL_CONSECUTIVO)
                        ):
                            repetida = (
                                previa,
                                "N° consecutivos, descripción idéntica y registro con un día de diferencia",
                            )
                    if repetida:
                        break
                if repetida:
                    break
        if repetida:
            retiradas.add(indice)
            previa = repetida[0]
            descartadas.append(
                {
                    "conservado": filas[previa].get(col_id) if col_id else previa + 1,
                    "descartado": fila.get(col_id) if col_id else indice + 1,
                    "motivo": repetida[1],
                }
            )
        elif dia is not None:
            por_dia.setdefault((grupo, dia.toordinal()), []).append(indice)
    conservadas = [fila for indice, fila in enumerate(filas) if indice not in retiradas]
    return conservadas, descartadas


def valor_calculado(calculada, fila, columnas_reales):
    origen = fila.get(columnas_reales[calculada["de"]])
    tipo = calculada["tipo"]
    if tipo in ("anio", "mes", "mes_nombre", "trimestre"):
        dia = fecha_de(origen)
        if dia is None:
            return None
        if tipo == "anio":
            return dia.year
        if tipo == "mes":
            return dia.month
        if tipo == "mes_nombre":
            return MESES[dia.month - 1]
        return f"T{(dia.month - 1) // 3 + 1}"
    if tipo == "clasificar":
        valor = sin_acentos(origen)
        for etiqueta, valores in (calculada.get("reglas") or {}).items():
            if valor in {sin_acentos(v) for v in (valores if isinstance(valores, list) else [valores])}:
                return etiqueta
        return calculada.get("otro", "Otro")
    raise ValueError(f"tipo de columna calculada '{tipo}' no válido. Usa: anio, mes, mes_nombre, trimestre, clasificar")


def _orden_bloque(columna, valores_conteo, calculadas):
    tipo = next((c["tipo"] for c in calculadas if c["nombre"] == columna), None)
    if tipo in ("anio", "mes", "trimestre"):
        return sorted(valores_conteo.items(), key=lambda par: str(par[0]))
    if tipo == "mes_nombre":
        return sorted(valores_conteo.items(), key=lambda par: MESES.index(par[0]) if par[0] in MESES else 99)
    return sorted(valores_conteo.items(), key=lambda par: (-par[1], str(par[0])))


def resumen_filas(filas, columnas_resumen, calculadas):
    bloques = []
    for columna in columnas_resumen:
        conteo = {}
        for fila in filas:
            valor = fila.get(columna)
            valor = "(sin dato)" if valor in (None, "") else valor
            conteo[valor] = conteo.get(valor, 0) + 1
        bloques.append((columna, _orden_bloque(columna, conteo, calculadas)))
    return bloques


def escribir_excel(ruta, hojas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    libro.remove(libro.active)
    relleno = PatternFill("solid", fgColor="0D1D35")
    for nombre, columnas, filas in hojas:
        hoja = libro.create_sheet(nombre[:31])
        hoja.append(columnas)
        for fila in filas:
            hoja.append([fila.get(c) if isinstance(fila, dict) else fila[i] for i, c in enumerate(columnas)])
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.freeze_panes = "A2"
        if filas:
            hoja.auto_filter.ref = hoja.dimensions
        for indice, columna in enumerate(columnas, 1):
            largo = max(
                [len(str(columna))]
                + [len(str(f.get(columna) if isinstance(f, dict) else f[indice - 1] or "")) for f in filas[:200]]
            )
            hoja.column_dimensions[get_column_letter(indice)].width = min(60, max(10, largo + 2))
    libro.save(ruta)

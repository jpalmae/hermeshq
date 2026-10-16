"""Genera una presentación usando los diseños (layouts) de una plantilla .pptx propia del usuario."""

from __future__ import annotations

import unicodedata
from pathlib import Path

EMU_POR_PULGADA = 914400
IMAGENES = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
TIPOS_TITULO = ("TITLE", "CENTER_TITLE", "VERTICAL_TITLE")
TIPOS_CUERPO = ("BODY", "OBJECT", "VERTICAL_BODY", "VERTICAL_OBJECT")
TIPOS_IGNORADOS = ("DATE", "FOOTER", "SLIDE_NUMBER", "HEADER")
MAX_LAMINAS = 60


class ErrorPlantilla(ValueError):
    pass


def _plano(texto) -> str:
    base = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(base.split())


def _tipo(ph) -> str:
    try:
        return ph.placeholder_format.type.name
    except (AttributeError, ValueError):
        return "DESCONOCIDO"


def _pulgadas(valor) -> float | None:
    return round(valor / EMU_POR_PULGADA, 2) if valor is not None else None


def abrir(ruta: Path):
    from pptx import Presentation

    if ruta.suffix.lower() != ".pptx":
        raise ErrorPlantilla(
            "La plantilla debe ser un archivo .pptx (si es .potx, guárdala como .pptx desde PowerPoint)."
        )
    if not ruta.is_file():
        raise ErrorPlantilla(f"No encuentro la plantilla: {ruta.name}")
    try:
        return Presentation(str(ruta))
    except Exception as exc:  # noqa: BLE001  # un archivo dañado no debe tumbar la tarea
        raise ErrorPlantilla(f"No pude abrir la plantilla: {str(exc)[:150]}") from exc


def inspeccionar(ruta: Path) -> dict:
    prs = abrir(ruta)
    disenos = []
    for indice, layout in enumerate(prs.slide_layouts):
        marcadores = []
        for ph in layout.placeholders:
            tipo = _tipo(ph)
            if tipo in TIPOS_IGNORADOS:
                continue
            marcadores.append(
                {
                    "idx": ph.placeholder_format.idx,
                    "tipo": tipo,
                    "nombre": ph.name,
                    "x": _pulgadas(ph.left),
                    "y": _pulgadas(ph.top),
                    "ancho": _pulgadas(ph.width),
                    "alto": _pulgadas(ph.height),
                }
            )
        disenos.append({"indice": indice, "nombre": layout.name, "marcadores": marcadores})
    laminas = []
    for numero, lamina in enumerate(prs.slides, 1):
        titulo = lamina.shapes.title.text if lamina.shapes.title is not None else ""
        laminas.append({"numero": numero, "diseno": lamina.slide_layout.name, "titulo": titulo[:80]})
    return {
        "ancho_pulgadas": _pulgadas(prs.slide_width),
        "alto_pulgadas": _pulgadas(prs.slide_height),
        "disenos": disenos,
        "laminas_existentes": laminas,
    }


def _quitar_laminas(prs) -> None:
    lista = prs.slides._sldIdLst
    for entrada in list(lista):
        prs.part.drop_rel(entrada.rId)
        lista.remove(entrada)


def _es_tipo(ph, tipos) -> bool:
    return _tipo(ph) in tipos


def _marcadores(layout, tipos):
    return [ph for ph in layout.placeholders if _es_tipo(ph, tipos)]


def _elegir_diseno(prs, indicado, spec: dict):
    disenos = list(prs.slide_layouts)
    if indicado is not None and indicado != "":
        if isinstance(indicado, int) or str(indicado).strip().isdigit():
            posicion = int(indicado)
            if not 0 <= posicion < len(disenos):
                raise ErrorPlantilla(f"El diseño {posicion} no existe (hay {len(disenos)}: 0 a {len(disenos) - 1}).")
            return disenos[posicion]
        buscado = _plano(indicado)
        exacto = [d for d in disenos if _plano(d.name) == buscado]
        parecido = [d for d in disenos if buscado in _plano(d.name)]
        elegido = (exacto or parecido or [None])[0]
        if elegido is None:
            nombres = ", ".join(d.name for d in disenos[:12])
            raise ErrorPlantilla(f"No encuentro el diseño '{indicado}'. Disponibles: {nombres}")
        return elegido
    quiere_imagen = bool(spec.get("image_path"))
    quiere_subtitulo = bool(spec.get("subtitle")) and not (spec.get("bullets") or spec.get("table"))
    for diseno in disenos:
        tipos = {_tipo(p) for p in diseno.placeholders}
        if quiere_subtitulo and "SUBTITLE" in tipos and tipos & set(TIPOS_TITULO):
            return diseno
    for diseno in disenos:
        tipos = {_tipo(p) for p in diseno.placeholders}
        if quiere_imagen and "PICTURE" in tipos and tipos & set(TIPOS_TITULO):
            return diseno
    for diseno in disenos:
        tipos = {_tipo(p) for p in diseno.placeholders}
        if tipos & set(TIPOS_TITULO) and tipos & set(TIPOS_CUERPO):
            return diseno
    return disenos[1] if len(disenos) > 1 else disenos[0]


def _escribir_parrafos(marco, elementos) -> None:
    marco.clear()
    primero = True
    for elemento in elementos:
        if isinstance(elemento, dict):
            texto, nivel = str(elemento.get("text") or elemento.get("texto") or ""), int(elemento.get("level") or 0)
        else:
            texto, nivel = str(elemento), 0
        parrafo = marco.paragraphs[0] if primero else marco.add_paragraph()
        parrafo.text = texto
        parrafo.level = max(0, min(8, nivel))
        primero = False


def _imagen(ruta_texto: str, base: Path) -> Path:
    ruta = Path(ruta_texto)
    if not ruta.is_absolute():
        ruta = base / ruta
    if not ruta.is_file():
        raise ErrorPlantilla(f"No encuentro la imagen: {ruta_texto}")
    if ruta.suffix.lower() not in IMAGENES:
        raise ErrorPlantilla(f"Formato de imagen no válido ({ruta.suffix}). Usa: {', '.join(sorted(IMAGENES))}")
    return ruta


def _colocar_imagen(prs, lamina, ruta: Path, area) -> None:
    izquierda, arriba, ancho, alto = area
    from PIL import Image

    with Image.open(ruta) as imagen:
        proporcion = imagen.width / max(1, imagen.height)
    alto_final = alto
    ancho_final = int(alto_final * proporcion)
    if ancho_final > ancho:
        ancho_final = ancho
        alto_final = int(ancho_final / proporcion)
    lamina.shapes.add_picture(
        str(ruta),
        izquierda + (ancho - ancho_final) // 2,
        arriba + (alto - alto_final) // 2,
        width=ancho_final,
        height=alto_final,
    )


def _area_libre(prs, lamina, cuerpo) -> tuple[int, int, int, int]:
    if cuerpo is not None:
        return cuerpo.left, cuerpo.top, cuerpo.width, cuerpo.height
    margen = int(0.5 * EMU_POR_PULGADA)
    titulo = lamina.shapes.title
    arriba = (titulo.top + titulo.height + margen // 2) if titulo is not None else margen
    return margen, arriba, prs.slide_width - 2 * margen, prs.slide_height - arriba - margen


def _poner_tabla(lamina, tabla: dict, area) -> None:
    encabezados = [str(c) for c in (tabla.get("headers") or tabla.get("encabezados") or [])]
    filas = [[str(c) for c in fila] for fila in (tabla.get("rows") or tabla.get("filas") or [])]
    if not encabezados and not filas:
        raise ErrorPlantilla("La tabla necesita headers y/o rows.")
    columnas = max([len(encabezados)] + [len(f) for f in filas])
    cuerpo = ([encabezados] if encabezados else []) + filas
    izquierda, arriba, ancho, alto = area
    forma = lamina.shapes.add_table(
        len(cuerpo), columnas, izquierda, arriba, ancho, min(alto, int(0.4 * EMU_POR_PULGADA) * len(cuerpo))
    )
    for r, fila in enumerate(cuerpo):
        for c in range(columnas):
            forma.table.cell(r, c).text = fila[c] if c < len(fila) else ""


def _buscar_marcador(layout_slide, clave):
    for ph in layout_slide.placeholders:
        if str(ph.placeholder_format.idx) == str(clave) or _plano(ph.name) == _plano(clave):
            return ph
    return None


def _quitar_marcador(ph) -> None:
    elemento = ph._element
    elemento.getparent().remove(elemento)


def construir(ruta_plantilla: Path, laminas: list, destino: Path, conservar: bool, base: Path) -> dict:
    if not isinstance(laminas, list) or not laminas:
        raise ErrorPlantilla("slides debe ser una lista con al menos una lámina.")
    if len(laminas) > MAX_LAMINAS:
        raise ErrorPlantilla(f"Máximo {MAX_LAMINAS} láminas por presentación.")
    prs = abrir(ruta_plantilla)
    existentes = len(prs.slides)
    if not conservar:
        _quitar_laminas(prs)
    avisos, usados = [], []
    for numero, spec in enumerate(laminas, 1):
        if not isinstance(spec, dict):
            raise ErrorPlantilla(f"La lámina {numero} debe ser un objeto.")
        try:
            diseno = _elegir_diseno(prs, spec.get("layout"), spec)
        except ErrorPlantilla as exc:
            raise ErrorPlantilla(f"Lámina {numero}: {exc}") from exc
        lamina = prs.slides.add_slide(diseno)
        usados.append(diseno.name)
        marcadores = list(lamina.placeholders)
        titulos = [p for p in marcadores if _es_tipo(p, TIPOS_TITULO)]
        subtitulos = [p for p in marcadores if _tipo(p) == "SUBTITLE"]
        cuerpos = [p for p in marcadores if _es_tipo(p, TIPOS_CUERPO)]
        fotos = [p for p in marcadores if _tipo(p) == "PICTURE"]
        if spec.get("title") and titulos:
            titulos[0].text_frame.text = str(spec["title"])
        elif spec.get("title"):
            avisos.append(f"Lámina {numero}: el diseño '{diseno.name}' no tiene marcador de título.")
        if spec.get("subtitle"):
            if subtitulos:
                subtitulos[0].text_frame.text = str(spec["subtitle"])
            elif cuerpos and not (spec.get("bullets") or spec.get("text")):
                cuerpos[0].text_frame.text = str(spec["subtitle"])
                cuerpos = cuerpos[1:]
            else:
                avisos.append(f"Lámina {numero}: el diseño '{diseno.name}' no tiene subtítulo; se omitió.")
        cuerpo_usado = None
        texto = spec.get("bullets") or ([spec["text"]] if spec.get("text") else [])
        if texto and cuerpos:
            cuerpo_usado = cuerpos[0]
            _escribir_parrafos(cuerpo_usado.text_frame, texto if isinstance(texto, list) else [texto])
        elif texto:
            avisos.append(f"Lámina {numero}: el diseño '{diseno.name}' no tiene marcador de texto; se omitió el texto.")
        for clave, valor in (spec.get("fields") or {}).items():
            ph = _buscar_marcador(lamina, clave)
            if ph is None or not ph.has_text_frame:
                avisos.append(f"Lámina {numero}: no existe el marcador '{clave}'.")
            else:
                ph.text_frame.text = str(valor)
        imagen = spec.get("image_path")
        if imagen:
            ruta = _imagen(str(imagen), base)
            if fotos:
                fotos[0].insert_picture(str(ruta))
                fotos = fotos[1:]
            else:
                libre = cuerpos[0] if (cuerpos and cuerpo_usado is None) else None
                if cuerpo_usado is not None:
                    cuerpo_usado.width = int(cuerpo_usado.width * 0.5)
                    area = (
                        cuerpo_usado.left + cuerpo_usado.width + int(0.3 * EMU_POR_PULGADA),
                        cuerpo_usado.top,
                        prs.slide_width - (cuerpo_usado.left + cuerpo_usado.width) - int(0.6 * EMU_POR_PULGADA),
                        cuerpo_usado.height,
                    )
                else:
                    area = _area_libre(prs, lamina, libre)
                    if libre is not None:
                        _quitar_marcador(libre)
                        cuerpos = cuerpos[1:]
                _colocar_imagen(prs, lamina, ruta, area)
        if spec.get("table"):
            libre = cuerpos[0] if (cuerpos and cuerpo_usado is None) else None
            area = _area_libre(prs, lamina, libre)
            if libre is not None:
                _quitar_marcador(libre)
            _poner_tabla(lamina, spec["table"], area)
        if spec.get("notes"):
            lamina.notes_slide.notes_text_frame.text = str(spec["notes"])
        for ph in list(lamina.placeholders):
            if _tipo(ph) in TIPOS_IGNORADOS:
                continue
            vacio = ph.has_text_frame and not ph.text_frame.text.strip()
            if vacio or (_tipo(ph) == "PICTURE" and not hasattr(ph, "image")):
                _quitar_marcador(ph)
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destino))
    return {
        "laminas": len(laminas) + (existentes if conservar else 0),
        "disenos_usados": usados,
        "avisos": avisos,
    }

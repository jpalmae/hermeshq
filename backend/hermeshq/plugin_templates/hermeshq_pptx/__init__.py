from __future__ import annotations

import json
import os
from pathlib import Path

_SUPPORTED_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}


def _check_requirements():
    try:
        import pptx  # noqa: F401

        return True
    except ImportError:
        return False


def _safe_filename(title: str, ext: str) -> str:
    safe_title = "".join(c if c.isalnum() or c in ("-", "_") else "-" for c in title.lower())[:60].strip("-")
    return f"{safe_title or 'presentacion'}.{ext}"


def _resolve_local_path(raw_path: str) -> Path:
    p = Path(raw_path)
    if not p.is_absolute():
        work_dir = Path(os.environ.get("HERMES_SESSION_CWD", os.getcwd()))
        p = work_dir / raw_path
    return p


def _resolve_slide_image(image_path: str) -> Path:
    resolved = _resolve_local_path(image_path)
    if not resolved.is_file():
        raise ValueError(f"image not found: {image_path}")
    if resolved.suffix.lower() not in _SUPPORTED_IMAGE_EXTS:
        raise ValueError(
            f"unsupported image format '{resolved.suffix}' for {image_path}. "
            f"Supported: {', '.join(sorted(_SUPPORTED_IMAGE_EXTS))}"
        )
    return resolved


def _add_slide_image(prs, slide, resolved_image: Path, width_inches: float | None, body_ph) -> None:
    from pptx.util import Inches

    if body_ph is not None:
        # Text is present — shrink its placeholder to the left half and put
        # the image in the remaining right half so they don't overlap.
        body_ph.width = int(body_ph.width * 0.5)
        left = body_ph.left + body_ph.width + Inches(0.3)
        top = body_ph.top
        max_width = prs.slide_width - left - Inches(0.3)
        width = Inches(width_inches) if width_inches else max_width
        width = min(width, max_width)
        slide.shapes.add_picture(str(resolved_image), left, top, width=int(width))
    else:
        # No text placeholder in use — center the image below the title.
        width = Inches(min(width_inches, 8.5)) if width_inches else Inches(6.0)
        width = min(width, prs.slide_width - Inches(1.0))
        left = int((prs.slide_width - width) / 2)
        top = Inches(1.8)
        slide.shapes.add_picture(str(resolved_image), left, top, width=int(width))


def _handle_generate_pptx(args, **_kwargs):
    title = str(args.get("title") or "Presentación").strip()
    subtitle = str(args.get("subtitle") or "").strip()
    slides = args.get("slides")
    filename = str(args.get("filename") or "").strip()

    try:
        from pptx import Presentation
    except ImportError:
        return json.dumps(
            {"success": False, "error": "python-pptx is not installed. PowerPoint generation is unavailable."}
        )

    if not isinstance(slides, list) or not slides:
        return json.dumps({"success": False, "error": "slides must be a non-empty list of slide objects"})

    if not filename:
        filename = _safe_filename(title, "pptx")
    elif not filename.lower().endswith(".pptx"):
        filename += ".pptx"

    work_dir = Path(os.environ.get("HERMES_SESSION_CWD", os.getcwd()))
    output_dir = work_dir / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    filepath = output_dir / filename

    try:
        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        if title_slide.shapes.title:
            title_slide.shapes.title.text = title
        if subtitle and len(title_slide.placeholders) > 1:
            title_slide.placeholders[1].text = subtitle

        content_layout = prs.slide_layouts[1]
        section_layout = prs.slide_layouts[2] if len(prs.slide_layouts) > 2 else content_layout

        for slide_spec in slides:
            if not isinstance(slide_spec, dict):
                continue
            layout = section_layout if str(slide_spec.get("layout", "")).strip() == "section" else content_layout
            slide = prs.slides.add_slide(layout)
            if slide.shapes.title:
                slide.shapes.title.text = str(slide_spec.get("title", ""))
            bullets = [str(b) for b in (slide_spec.get("bullets") or [])]
            body_ph = None
            if bullets and len(slide.placeholders) > 1:
                body_ph = slide.placeholders[1]
                text_frame = body_ph.text_frame
                text_frame.clear()
                text_frame.text = bullets[0]
                for bullet in bullets[1:]:
                    paragraph = text_frame.add_paragraph()
                    paragraph.text = bullet
            image_path = str(slide_spec.get("image_path") or "").strip()
            if image_path:
                resolved_image = _resolve_slide_image(image_path)
                width_inches = slide_spec.get("image_width_inches")
                try:
                    width_inches = float(width_inches) if width_inches is not None else None
                except (TypeError, ValueError):
                    width_inches = None
                _add_slide_image(prs, slide, resolved_image, width_inches, body_ph)
            notes = str(slide_spec.get("notes") or "").strip()
            if notes:
                slide.notes_slide.notes_text_frame.text = notes

        prs.save(str(filepath))
    except Exception as exc:
        return json.dumps({"success": False, "error": f"PowerPoint generation failed: {exc}"})

    return json.dumps(
        {
            "success": True,
            "filename": filename,
            "path": str(filepath),
            "size_bytes": filepath.stat().st_size,
            "slide_count": len(slides) + 1,
        }
    )


def register(ctx):
    spec = {
        "name": "hq_pptx_generate",
        "description": (
            "Generate a PowerPoint (.pptx) presentation with a title slide plus content slides. "
            "Each content slide has a title, optional bullet points, an optional embedded image, and optional "
            "speaker notes. The file is saved to the agent's output directory and can be attached to messages. "
            "ALWAYS use this tool for PowerPoint files instead of python-pptx scripts with execute_code."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "title": {
                    "type": "string",
                    "description": "Presentation title, shown on the title slide and used for the default filename",
                },
                "subtitle": {
                    "type": "string",
                    "description": "Optional subtitle shown on the title slide",
                },
                "slides": {
                    "type": "array",
                    "description": "Ordered list of content slides (the title slide is added automatically before these)",
                    "items": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string", "description": "Slide title"},
                            "bullets": {
                                "type": "array",
                                "items": {"type": "string"},
                                "description": "Bullet points for the slide body",
                            },
                            "notes": {
                                "type": "string",
                                "description": "Optional speaker notes for this slide",
                            },
                            "image_path": {
                                "type": "string",
                                "description": (
                                    "Path to an existing local image file (png/jpg/jpeg/gif/bmp), relative to the "
                                    "workspace or absolute. The file must already exist (e.g. uploaded by the user "
                                    "or produced by another tool); this does not fetch URLs. If the slide also has "
                                    "bullets, the image is placed beside them; otherwise it's centered below the title."
                                ),
                            },
                            "image_width_inches": {
                                "type": "number",
                                "description": "Optional image display width in inches. If omitted, a sensible default is used based on the layout.",
                            },
                            "layout": {
                                "type": "string",
                                "enum": ["content", "section"],
                                "description": "'content' (default) for a normal title+bullets slide, 'section' for a section-divider slide",
                            },
                        },
                        "required": ["title"],
                    },
                },
                "filename": {
                    "type": "string",
                    "description": "Optional output filename (e.g. 'propuesta.pptx'). If omitted, derived from title.",
                },
            },
            "required": ["title", "slides"],
        },
        "handler": _handle_generate_pptx,
        "emoji": "📊",
    }

    ctx.register_tool(
        name=spec["name"],
        toolset="hermeshq_pptx",
        schema={
            "name": spec["name"],
            "description": spec["description"],
            "parameters": spec["parameters"],
        },
        handler=spec["handler"],
        check_fn=_check_requirements,
        description=spec["description"],
        emoji=spec["emoji"],
    )

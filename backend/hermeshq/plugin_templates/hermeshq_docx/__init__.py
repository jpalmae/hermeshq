from __future__ import annotations

import json
import os
from pathlib import Path

_BLOCK_TYPES = {
    "heading1",
    "heading2",
    "heading3",
    "paragraph",
    "bullet_list",
    "numbered_list",
    "table",
    "image",
    "page_break",
}

_SUPPORTED_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}


def _check_requirements():
    try:
        import docx  # noqa: F401

        return True
    except ImportError:
        return False


def _safe_filename(title: str, ext: str) -> str:
    safe_title = "".join(c if c.isalnum() or c in ("-", "_") else "-" for c in title.lower())[:60].strip("-")
    return f"{safe_title or 'documento'}.{ext}"


def _resolve_local_path(raw_path: str) -> Path:
    p = Path(raw_path)
    if not p.is_absolute():
        work_dir = Path(os.environ.get("HERMES_SESSION_CWD", os.getcwd()))
        p = work_dir / raw_path
    return p


def _add_image(doc, block: dict) -> None:
    from docx.shared import Inches

    raw_path = str(block.get("path") or "").strip()
    if not raw_path:
        raise ValueError("image block is missing 'path'")
    resolved = _resolve_local_path(raw_path)
    if not resolved.is_file():
        raise ValueError(f"image not found: {raw_path}")
    if resolved.suffix.lower() not in _SUPPORTED_IMAGE_EXTS:
        raise ValueError(
            f"unsupported image format '{resolved.suffix}' for {raw_path}. "
            f"Supported: {', '.join(sorted(_SUPPORTED_IMAGE_EXTS))}"
        )
    width_inches = block.get("width_inches")
    try:
        width_inches = float(width_inches) if width_inches is not None else 6.0
    except (TypeError, ValueError):
        width_inches = 6.0
    width_inches = max(0.5, min(width_inches, 6.5))
    doc.add_picture(str(resolved), width=Inches(width_inches))
    doc.paragraphs[-1].alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER
    caption = str(block.get("caption") or "").strip()
    if caption:
        caption_para = doc.add_paragraph(caption)
        caption_para.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER
        for run in caption_para.runs:
            run.italic = True


def _add_table(doc, block: dict) -> None:
    headers = [str(h) for h in (block.get("headers") or [])]
    rows = [r for r in (block.get("rows") or []) if isinstance(r, list)]
    ncols = len(headers) if headers else (len(rows[0]) if rows else 0)
    if not ncols:
        return
    total_rows = (1 if headers else 0) + len(rows)
    if not total_rows:
        return
    table = doc.add_table(rows=total_rows, cols=ncols)
    table.style = "Light Grid Accent 1"
    r = 0
    if headers:
        for i, h in enumerate(headers):
            table.cell(r, i).text = h
        r += 1
    for row in rows:
        for i in range(ncols):
            table.cell(r, i).text = str(row[i]) if i < len(row) else ""
        r += 1


def _handle_generate_docx(args, **_kwargs):
    title = str(args.get("title") or "Documento").strip()
    sections = args.get("sections")
    filename = str(args.get("filename") or "").strip()

    try:
        from docx import Document
    except ImportError:
        return json.dumps(
            {"success": False, "error": "python-docx is not installed. Word document generation is unavailable."}
        )

    if not isinstance(sections, list) or not sections:
        return json.dumps({"success": False, "error": "sections must be a non-empty list of content blocks"})

    if not filename:
        filename = _safe_filename(title, "docx")
    elif not filename.lower().endswith(".docx"):
        filename += ".docx"

    work_dir = Path(os.environ.get("HERMES_SESSION_CWD", os.getcwd()))
    output_dir = work_dir / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    filepath = output_dir / filename

    try:
        doc = Document()
        doc.add_heading(title, level=0)
        for block in sections:
            if not isinstance(block, dict):
                continue
            btype = str(block.get("type", "")).strip()
            if btype not in _BLOCK_TYPES:
                continue
            if btype in ("heading1", "heading2", "heading3"):
                doc.add_heading(str(block.get("text", "")), level=int(btype[-1]))
            elif btype == "paragraph":
                doc.add_paragraph(str(block.get("text", "")))
            elif btype == "bullet_list":
                for item in block.get("items") or []:
                    doc.add_paragraph(str(item), style="List Bullet")
            elif btype == "numbered_list":
                for item in block.get("items") or []:
                    doc.add_paragraph(str(item), style="List Number")
            elif btype == "table":
                _add_table(doc, block)
            elif btype == "image":
                _add_image(doc, block)
            elif btype == "page_break":
                doc.add_page_break()
        doc.save(str(filepath))
    except Exception as exc:
        return json.dumps({"success": False, "error": f"Word document generation failed: {exc}"})

    return json.dumps(
        {
            "success": True,
            "filename": filename,
            "path": str(filepath),
            "size_bytes": filepath.stat().st_size,
        }
    )


def register(ctx):
    spec = {
        "name": "hq_docx_generate",
        "description": (
            "Generate a Word (.docx) document from structured content blocks. "
            "Supports headings, paragraphs, bullet/numbered lists, tables, embedded images, and page breaks. "
            "The file is saved to the agent's output directory and can be attached to messages. "
            "ALWAYS use this tool for Word files instead of python-docx scripts with execute_code."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "title": {
                    "type": "string",
                    "description": "Document title (shown as the top-level heading and used for the default filename)",
                },
                "sections": {
                    "type": "array",
                    "description": "Ordered list of content blocks that make up the document body",
                    "items": {
                        "type": "object",
                        "properties": {
                            "type": {
                                "type": "string",
                                "enum": sorted(_BLOCK_TYPES),
                                "description": "Block type",
                            },
                            "text": {
                                "type": "string",
                                "description": "Text content — used by heading1/heading2/heading3/paragraph blocks",
                            },
                            "items": {
                                "type": "array",
                                "items": {"type": "string"},
                                "description": "List items — used by bullet_list/numbered_list blocks",
                            },
                            "headers": {
                                "type": "array",
                                "items": {"type": "string"},
                                "description": "Optional table column headers — used by table blocks",
                            },
                            "rows": {
                                "type": "array",
                                "items": {"type": "array", "items": {"type": "string"}},
                                "description": "Table rows, each a list of cell values — used by table blocks",
                            },
                            "path": {
                                "type": "string",
                                "description": (
                                    "Path to an existing local image file (png/jpg/jpeg/gif/bmp), relative to the "
                                    "workspace or absolute — used by image blocks. The file must already exist "
                                    "(e.g. uploaded by the user or produced by another tool); this does not fetch URLs."
                                ),
                            },
                            "width_inches": {
                                "type": "number",
                                "description": "Optional image display width in inches (0.5–6.5, default 6.0) — used by image blocks",
                            },
                            "caption": {
                                "type": "string",
                                "description": "Optional italic caption shown centered below the image — used by image blocks",
                            },
                        },
                        "required": ["type"],
                    },
                },
                "filename": {
                    "type": "string",
                    "description": "Optional output filename (e.g. 'informe.docx'). If omitted, derived from title.",
                },
            },
            "required": ["title", "sections"],
        },
        "handler": _handle_generate_docx,
        "emoji": "📝",
    }

    ctx.register_tool(
        name=spec["name"],
        toolset="hermeshq_docx",
        schema={
            "name": spec["name"],
            "description": spec["description"],
            "parameters": spec["parameters"],
        },
        handler=spec["handler"],
        check_fn=_check_requirements,
        description=spec["description"],
        emoji=spec["emoji"],
    )

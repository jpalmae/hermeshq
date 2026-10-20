from __future__ import annotations

import json
import os
import re
from pathlib import Path

_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
_SUPPORTED_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}


def _check_requirements():
    try:
        import openpyxl  # noqa: F401

        return True
    except ImportError:
        return False


def _safe_filename(title: str, ext: str) -> str:
    safe_title = "".join(c if c.isalnum() or c in ("-", "_") else "-" for c in title.lower())[:60].strip("-")
    return f"{safe_title or 'hoja-de-calculo'}.{ext}"


def _sanitize_sheet_name(name: str, index: int, taken: set[str]) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub(" ", name or "").strip()[:31] or f"Hoja{index}"
    candidate = cleaned
    suffix = 2
    while candidate.lower() in taken:
        candidate = f"{cleaned[:28]} {suffix}"
        suffix += 1
    taken.add(candidate.lower())
    return candidate


def _resolve_local_path(raw_path: str) -> Path:
    p = Path(raw_path)
    if not p.is_absolute():
        work_dir = Path(os.environ.get("HERMES_SESSION_CWD", os.getcwd()))
        p = work_dir / raw_path
    return p


def _add_sheet_images(ws, images: list, ncols: int, get_column_letter) -> None:
    from openpyxl.drawing.image import Image as XLImage

    default_col = ncols + 2 if ncols else 2
    for idx, img_spec in enumerate(images):
        if not isinstance(img_spec, dict):
            continue
        raw_path = str(img_spec.get("path") or "").strip()
        if not raw_path:
            raise ValueError("image entry is missing 'path'")
        resolved = _resolve_local_path(raw_path)
        if not resolved.is_file():
            raise ValueError(f"image not found: {raw_path}")
        if resolved.suffix.lower() not in _SUPPORTED_IMAGE_EXTS:
            raise ValueError(
                f"unsupported image format '{resolved.suffix}' for {raw_path}. "
                f"Supported: {', '.join(sorted(_SUPPORTED_IMAGE_EXTS))}"
            )
        xl_image = XLImage(str(resolved))
        width_inches = img_spec.get("width_inches")
        try:
            width_inches = float(width_inches) if width_inches is not None else None
        except (TypeError, ValueError):
            width_inches = None
        if width_inches and xl_image.width:
            width_inches = max(0.5, min(width_inches, 10.0))
            scale = (width_inches * 96) / xl_image.width
            xl_image.width = int(xl_image.width * scale)
            xl_image.height = int(xl_image.height * scale)
        anchor = str(img_spec.get("anchor") or "").strip()
        if not anchor:
            anchor = f"{get_column_letter(default_col)}{1 + idx * 16}"
        ws.add_image(xl_image, anchor)


def _handle_generate_xlsx(args, **_kwargs):
    title = str(args.get("title") or "Hoja de cálculo").strip()
    sheets = args.get("sheets")
    filename = str(args.get("filename") or "").strip()

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return json.dumps(
            {"success": False, "error": "openpyxl is not installed. Excel spreadsheet generation is unavailable."}
        )

    if not isinstance(sheets, list) or not sheets:
        return json.dumps({"success": False, "error": "sheets must be a non-empty list of sheet objects"})

    if not filename:
        filename = _safe_filename(title, "xlsx")
    elif not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"

    work_dir = Path(os.environ.get("HERMES_SESSION_CWD", os.getcwd()))
    output_dir = work_dir / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    filepath = output_dir / filename

    try:
        wb = Workbook()
        wb.remove(wb.active)
        taken_names: set[str] = set()
        for i, sheet_spec in enumerate(sheets, start=1):
            if not isinstance(sheet_spec, dict):
                continue
            name = _sanitize_sheet_name(str(sheet_spec.get("name") or ""), i, taken_names)
            ws = wb.create_sheet(title=name)
            headers = [str(h) for h in (sheet_spec.get("headers") or [])]
            rows = [r for r in (sheet_spec.get("rows") or []) if isinstance(r, list)]
            ncols = len(headers) if headers else (max((len(r) for r in rows), default=0))

            r = 1
            if headers:
                for c, h in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=c, value=h)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="0F3460", end_color="0F3460", fill_type="solid")
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                ws.freeze_panes = "A2"
                r = 2

            for row in rows:
                for c, value in enumerate(row, start=1):
                    ws.cell(row=r, column=c, value=value)
                r += 1

            for c in range(1, ncols + 1):
                max_len = len(str(headers[c - 1])) if headers and c - 1 < len(headers) else 0
                for row in rows:
                    if c - 1 < len(row) and row[c - 1] is not None:
                        max_len = max(max_len, len(str(row[c - 1])))
                ws.column_dimensions[get_column_letter(c)].width = min(max(max_len + 2, 8), 50)

            images = sheet_spec.get("images") or []
            if images:
                _add_sheet_images(ws, images, ncols, get_column_letter)

        if not wb.sheetnames:
            wb.create_sheet(title="Hoja1")
        wb.save(str(filepath))
    except Exception as exc:
        return json.dumps({"success": False, "error": f"Excel spreadsheet generation failed: {exc}"})

    return json.dumps(
        {
            "success": True,
            "filename": filename,
            "path": str(filepath),
            "size_bytes": filepath.stat().st_size,
            "sheet_count": len(wb.sheetnames),
        }
    )


def register(ctx):
    spec = {
        "name": "hq_xlsx_generate",
        "description": (
            "Generate an Excel (.xlsx) workbook with one or more sheets, each from an optional header row, "
            "data rows, and optional embedded images. Headers are styled bold with a colored background and "
            "the header row is frozen. The file is saved to the agent's output directory and can be attached "
            "to messages. ALWAYS use this tool for Excel files instead of openpyxl scripts with execute_code."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "title": {
                    "type": "string",
                    "description": "Workbook title, used for the default filename",
                },
                "sheets": {
                    "type": "array",
                    "description": "Ordered list of sheets to create",
                    "items": {
                        "type": "object",
                        "properties": {
                            "name": {
                                "type": "string",
                                "description": "Sheet/tab name (max 31 chars, invalid characters are stripped)",
                            },
                            "headers": {
                                "type": "array",
                                "items": {"type": "string"},
                                "description": "Optional column headers for the first row",
                            },
                            "rows": {
                                "type": "array",
                                "items": {
                                    "type": "array",
                                    "items": {"type": ["string", "number", "boolean", "null"]},
                                },
                                "description": "Data rows, each a list of cell values (string or number)",
                            },
                            "images": {
                                "type": "array",
                                "description": "Optional images to embed in this sheet",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "path": {
                                            "type": "string",
                                            "description": (
                                                "Path to an existing local image file (png/jpg/jpeg/gif/bmp), "
                                                "relative to the workspace or absolute. The file must already "
                                                "exist (e.g. uploaded by the user or produced by another tool); "
                                                "this does not fetch URLs."
                                            ),
                                        },
                                        "anchor": {
                                            "type": "string",
                                            "description": "Optional top-left anchor cell (e.g. 'F2'). If omitted, placed to the right of the data.",
                                        },
                                        "width_inches": {
                                            "type": "number",
                                            "description": "Optional display width in inches (aspect ratio is preserved). If omitted, the image's native size is used.",
                                        },
                                    },
                                    "required": ["path"],
                                },
                            },
                        },
                        "required": ["rows"],
                    },
                },
                "filename": {
                    "type": "string",
                    "description": "Optional output filename (e.g. 'ventas.xlsx'). If omitted, derived from title.",
                },
            },
            "required": ["title", "sheets"],
        },
        "handler": _handle_generate_xlsx,
        "emoji": "📈",
    }

    ctx.register_tool(
        name=spec["name"],
        toolset="hermeshq_xlsx",
        schema={
            "name": spec["name"],
            "description": spec["description"],
            "parameters": spec["parameters"],
        },
        handler=spec["handler"],
        check_fn=_check_requirements,
        description=spec["description"],
        emoji=spec["emoji"],
    )

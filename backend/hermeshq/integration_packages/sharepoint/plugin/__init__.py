from __future__ import annotations

import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import urllib.error
import urllib.parse
import urllib.request
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

GRAPH_BASE = "https://graph.microsoft.com/v1.0"
TOOLSET = "hermeshq_sharepoint"

_MAX_DOWNLOAD_BYTES = 10 * 1024 * 1024
_DOWNLOAD_DIR = "sharepoint_downloads"
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
_DEFAULT_MAX_CHARS = 20_000
_HARD_MAX_CHARS = 100_000
_MAX_PDF_PAGES = 200
_MAX_SHEET_ROWS = 2000
_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".yaml",
    ".yml",
    ".ini",
    ".rtf",
}


def _resolve_session_user_id() -> str | None:
    """Resolve the HermesHQ user for the in-flight live gateway turn.

    Hermes Agent tracks the platform/sender of the current turn per-message
    via contextvars (safe under concurrent messages in the same gateway
    process), exposed through these env vars. Unused for punctual tasks,
    which carry their user in HERMESHQ_TASK_PAYLOAD instead.
    """
    platform = os.environ.get("HERMES_SESSION_PLATFORM", "").strip().lower()
    sender_id = os.environ.get("HERMES_SESSION_USER_ID", "").strip()
    if not platform or not sender_id:
        return None
    base_url = os.environ.get("HERMESHQ_INTERNAL_API_URL", "").rstrip("/")
    agent_id = os.environ.get("HERMESHQ_AGENT_ID", "")
    agent_token = os.environ.get("HERMESHQ_AGENT_TOKEN", "")
    if not base_url or not agent_id or not agent_token:
        return None
    url = (
        f"{base_url}/control/resolve-channel-user"
        f"?platform={urllib.parse.quote(platform)}&sender_id={urllib.parse.quote(sender_id)}"
    )
    req = urllib.request.Request(
        url,
        method="GET",
        headers={
            "X-HermesHQ-Agent-ID": agent_id,
            "X-HermesHQ-Agent-Token": agent_token,
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        return str(data.get("hermeshq_user_id") or "").strip() or None
    except urllib.error.HTTPError:
        return None
    except Exception:  # noqa: BLE001  # HTTP request catch-all
        return None


def _task_user_id() -> str | None:
    raw = os.environ.get("HERMESHQ_TASK_PAYLOAD", "")
    if raw:
        try:
            payload = json.loads(raw)
            meta = payload.get("metadata") or {}
            uid = str(meta.get("thread_user_id") or meta.get("created_by_user_id") or "").strip()
            if uid:
                return uid
        except (json.JSONDecodeError, ValueError):
            pass
    resolved = _resolve_session_user_id()
    if resolved:
        return resolved
    return os.environ.get("HERMESHQ_RESOLVED_USER_ID") or None


def _get_m365_token(user_id: str) -> tuple[str | None, str | None, str]:
    """Returns (access_token, sharepoint_site_url, error_detail)."""
    base_url = os.environ.get("HERMESHQ_INTERNAL_API_URL", "").rstrip("/")
    agent_id = os.environ.get("HERMESHQ_AGENT_ID", "")
    agent_token = os.environ.get("HERMESHQ_AGENT_TOKEN", "")
    if not base_url or not agent_id or not agent_token:
        return None, None, "HermesHQ internal control no configurado"
    url = f"{base_url}/control/m365/agent-token?user_id={user_id}"
    req = urllib.request.Request(
        url,
        method="GET",
        headers={"X-HermesHQ-Agent-ID": agent_id, "X-HermesHQ-Agent-Token": agent_token},
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            site_url = (data.get("sharepoint_site_url") or "").strip().rstrip("/") or None
            return data.get("access_token"), site_url, ""
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        try:
            detail = json.loads(body).get("detail", body)
        except (json.JSONDecodeError, ValueError):
            detail = body
        return None, None, str(detail)
    except Exception as exc:  # noqa: BLE001  # HTTP request catch-all
        return None, None, str(exc)


def _encode_url(url: str) -> str:
    return urllib.parse.quote(url, safe="%:/?#[]@!$&'()*+,;=~")


def _graph(method: str, path: str, access_token: str, payload: dict | None = None) -> dict:
    url = _encode_url(path if path.startswith("https://") else f"{GRAPH_BASE}{path}")
    data = json.dumps(payload).encode("utf-8") if payload else None
    req = urllib.request.Request(
        url,
        data=data,
        method=method.upper(),
        headers={"Authorization": f"Bearer {access_token}", "Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            raw = resp.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        retry_after = exc.headers.get("Retry-After") if exc.headers else None
        try:
            parsed = json.loads(body)
            err = parsed.get("error", {})
            code = err.get("code", "")
            msg = err.get("message", body)
        except (json.JSONDecodeError, ValueError):
            code, msg = "", body
        _FRIENDLY = {
            "Authorization_RequestDenied": "Sin permiso para acceder a este recurso (403).",
            "accessDenied": "Sin permiso para acceder a este recurso (403).",
            "itemNotFound": "Archivo o carpeta no encontrado (404).",
            "ActivityLimitReached": f"Límite de peticiones alcanzado. Reintenta en {retry_after or '60'} segundos.",
        }
        friendly = _FRIENDLY.get(code) or (f"Error {exc.code}: {msg}" if msg else f"Error HTTP {exc.code}")
        result: dict = {"error": friendly, "_graph_code": code}
        if retry_after:
            result["retry_after"] = retry_after
        return result


def _auth_error(detail: str) -> str:
    return json.dumps(
        {
            "success": False,
            "error": f"No se pudo obtener token M365: {detail}. Verifica que el usuario haya conectado su cuenta Microsoft 365 en Mi cuenta.",
        }
    )


# ── Tool handlers ─────────────────────────────────────────────────────────────


_LOOKUP_HINT = (
    " Pista: no adivines rutas, nombres de sitio ni ids. Busca primero con sharepoint_search (palabras del nombre) y "
    "usa el item_id y drive_id del resultado con sharepoint_read_file o sharepoint_download_file; para ver los sitios "
    "y bibliotecas disponibles usa sharepoint_list_drives."
)
_LOOKUP_ERROR_MARKERS = ("404", "root:", "Url specified is invalid", "Invalid request", "no encontrado")


def _with_hint(error: str) -> str:
    text = str(error)
    if any(marker in text for marker in _LOOKUP_ERROR_MARKERS) and "Pista:" not in text:
        return text + _LOOKUP_HINT
    return text


_MAX_TREE_ITEMS = 300


def _fetch_children(base: str, folder: str, token: str) -> tuple[list | None, str | None]:
    path: str | None = f"{base}:/{folder}:/children" if folder else f"{base}/children"
    items: list = []
    while path and len(items) < 200:
        result = _graph("GET", path, token)
        if "error" in result:
            return None, result["error"]
        items.extend(result.get("value", []))
        next_link: str = result.get("@odata.nextLink") or ""
        path = next_link if (next_link and len(items) < 200) else None
    return items, None


def _site_bases(site_url: str, folder_path: str, token: str) -> list[tuple[str, str, str | None]]:
    """Candidatos (base, carpeta, nombre de biblioteca): prueba la biblioteca nombrada en la ruta y luego todas las del sitio."""
    parsed = urllib.parse.urlparse(site_url)
    hostname = parsed.netloc
    site_path = parsed.path.rstrip("/") or "/"
    default = (f"/sites/{hostname}:{site_path}:/drive/root", folder_path, None)
    site = _graph("GET", f"/sites/{hostname}:{site_path}?$select=id", token)
    if "error" in site or not site.get("id"):
        return [default]
    drives = _graph("GET", f"/sites/{site['id']}/drives?$select=id,name", token).get("value", [])
    partes = [p for p in folder_path.split("/") if p]
    candidatos: list[tuple[str, str, str | None]] = []
    if partes:
        for drive in drives:
            if _fold_name(drive.get("name")) == _fold_name(partes[0]):
                candidatos.append((f"/drives/{drive['id']}/root", "/".join(partes[1:]), drive.get("name")))
    candidatos.append(default)
    for drive in drives:
        candidatos.append((f"/drives/{drive['id']}/root", folder_path, drive.get("name")))
    return candidatos


def _fold_name(value) -> str:
    import unicodedata

    texto = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "", texto.lower())


def _simplify_item(i: dict) -> dict:
    return {
        "id": i.get("id"),
        "name": i.get("name"),
        "type": "folder" if "folder" in i else "file",
        "size": i.get("size"),
        "url": i.get("webUrl"),
        "modified": i.get("lastModifiedDateTime"),
    }


def _build_tree(base: str, folder: str, token: str, depth: int, budget: list[int]) -> list[dict]:
    items, error = _fetch_children(base, folder, token)
    if items is None:
        return [{"name": f"(no se pudo leer: {str(error)[:60]})", "type": "error"}]
    nodes: list[dict] = []
    for item in items:
        if budget[0] <= 0:
            nodes.append({"name": "… (se omitió el resto por tamaño)", "type": "nota"})
            break
        budget[0] -= 1
        node = {"name": item.get("name"), "type": "folder" if "folder" in item else "file"}
        if node["type"] == "file":
            node["size"] = item.get("size")
        elif depth > 1:
            sub = f"{folder}/{item.get('name')}" if folder else str(item.get("name"))
            node["children"] = _build_tree(base, sub, token, depth - 1, budget)
        nodes.append(node)
    return nodes


def _build_folder_tree(
    base: str, folder: str, token: str, depth: int, budget: list[int]
) -> tuple[list[dict], int | None]:
    items, error = _fetch_children(base, folder, token)
    if items is None:
        return [{"name": f"(no se pudo leer: {str(error)[:60]})", "type": "error"}], None
    nodes: list[dict] = []
    for item in items:
        if "folder" not in item:
            continue
        if budget[0] <= 0:
            nodes.append(
                {
                    "name": "… (se omitieron más carpetas por tamaño; pide una subcarpeta o sube max_items)",
                    "type": "nota",
                }
            )
            break
        budget[0] -= 1
        node = {"name": item.get("name"), "type": "folder"}
        if depth > 1:
            sub = f"{folder}/{item.get('name')}" if folder else str(item.get("name"))
            node["children"], archivos = _build_folder_tree(base, sub, token, depth - 1, budget)
            if archivos is not None:
                node["archivos"] = archivos
        nodes.append(node)
    return nodes, sum(1 for i in items if "folder" not in i)


def _list_files_tool(args: dict, **_kwargs) -> str:
    """Lista archivos con Files.Read.All en OneDrive o SharePoint; depth 2-3 devuelve el árbol de subcarpetas en una sola llamada."""
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."})
    token, user_site_url, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err)

    site_url = str(args.get("site_url") or user_site_url or "").strip().rstrip("/")
    folder_path = str(args.get("folder_path") or "").strip().strip("/")
    solo_carpetas = bool(args.get("solo_carpetas"))
    try:
        depth = max(1, min(4 if solo_carpetas else 3, int(args.get("depth") or 1)))
    except (TypeError, ValueError):
        depth = 1
    try:
        max_items = max(50, min(800, int(args.get("max_items") or _MAX_TREE_ITEMS)))
    except (TypeError, ValueError):
        max_items = _MAX_TREE_ITEMS
    if not site_url and folder_path:
        folder_path, folder_error = _normalize_drive_folder(folder_path)
        if folder_error:
            return json.dumps({"success": False, "error": folder_error}, ensure_ascii=False)

    candidatos = _site_bases(site_url, folder_path, token) if site_url else [("/me/drive/root", folder_path, None)]
    raw_items: list | None = None
    last_error = "No se pudo listar la carpeta."
    usado: tuple[str, str, str | None] | None = None
    for base, carpeta, nombre in candidatos:
        raw_items, error = _fetch_children(base, carpeta, token)
        if raw_items is not None:
            usado = (base, carpeta, nombre)
            break
        last_error = error or last_error
    if raw_items is None or usado is None:
        return json.dumps({"success": False, "error": _with_hint(last_error)}, ensure_ascii=False)
    base, carpeta, nombre = usado
    resultado: dict = {"success": True}
    if nombre:
        resultado["biblioteca"] = nombre
    if depth > 1:
        presupuesto = [max_items]
        if solo_carpetas:
            arbol, archivos_raiz = _build_folder_tree(base, carpeta, token, depth, presupuesto)
            resultado["tree"] = arbol
            resultado["archivos_en_la_raiz"] = archivos_raiz
        else:
            resultado["tree"] = _build_tree(base, carpeta, token, depth, presupuesto)
        resultado["count"] = max_items - presupuesto[0]
        return json.dumps(resultado, ensure_ascii=False)
    simplified = [_simplify_item(i) for i in raw_items]
    resultado.update({"count": len(simplified), "items": simplified})
    return json.dumps(resultado, ensure_ascii=False)


def _get_file_tool(args: dict, **_kwargs) -> str:
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."})
    token, user_site_url, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err)

    file_path = str(args.get("file_path") or "").strip().strip("/")
    site_url = str(args.get("site_url") or user_site_url or "").strip().rstrip("/")

    if not file_path:
        return json.dumps({"success": False, "error": "Se requiere file_path."})

    if site_url:
        parsed = urllib.parse.urlparse(site_url)
        hostname = parsed.netloc
        site_path = parsed.path.rstrip("/") or "/"
        path = f"/sites/{hostname}:{site_path}:/drive/root:/{file_path}"
    else:
        path = f"/me/drive/root:/{file_path}"

    result = _graph("GET", path, token)
    if "error" in result:
        return json.dumps({"success": False, "error": _with_hint(result["error"])}, ensure_ascii=False)
    return json.dumps({"success": True, "item": result}, ensure_ascii=False)


def _list_drives_tool(args: dict, **_kwargs) -> str:
    """List accessible drives (OneDrive + SharePoint document libraries)."""
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."})
    token, _, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err)

    path: str | None = "/me/drives"
    raw_drives: list = []
    while path and len(raw_drives) < 200:
        result = _graph("GET", path, token)
        if "error" in result:
            return json.dumps({"success": False, "error": result["error"]})
        raw_drives.extend(result.get("value", []))
        next_link: str = result.get("@odata.nextLink") or ""
        path = next_link if (next_link and len(raw_drives) < 200) else None

    simplified = [
        {"id": d.get("id"), "name": d.get("name"), "type": d.get("driveType"), "url": d.get("webUrl")}
        for d in raw_drives
    ]
    return json.dumps({"success": True, "count": len(simplified), "drives": simplified}, ensure_ascii=False)


def _search_tool(args: dict, **_kwargs) -> str:
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."})
    token, _, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err)

    query = str(args.get("query") or "").strip()
    if not query:
        return json.dumps({"success": False, "error": "Se requiere query."})
    count = min(int(args.get("count") or 10), 25)

    payload = {
        "requests": [
            {
                "entityTypes": ["driveItem"],
                "query": {"queryString": query},
                "from": 0,
                "size": count,
                "fields": ["id", "name", "webUrl", "lastModifiedDateTime", "size", "parentReference"],
            }
        ]
    }
    result = _graph("POST", "/search/query", token, payload)
    if "error" in result:
        return json.dumps({"success": False, "error": result["error"]})
    hits = []
    for resp in result.get("value", []):
        for hc in resp.get("hitsContainers", []):
            for hit in hc.get("hits", []):
                resource = hit.get("resource", {})
                hits.append(
                    {
                        "name": resource.get("name"),
                        "url": resource.get("webUrl"),
                        "modified": resource.get("lastModifiedDateTime"),
                        "size": resource.get("size"),
                        "item_id": resource.get("id"),
                        "drive_id": (resource.get("parentReference") or {}).get("driveId"),
                    }
                )
    return json.dumps({"success": True, "count": len(hits), "hits": hits}, ensure_ascii=False)


# ── Plugin registration ───────────────────────────────────────────────────────


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):  # noqa: D401
        return None


def _download_bytes(url: str, access_token: str | None = None) -> tuple[bytes | None, str]:
    headers = {"Authorization": f"Bearer {access_token}"} if access_token else {}
    req = urllib.request.Request(_encode_url(url), headers=headers, method="GET")
    opener = urllib.request.build_opener(_NoRedirect)
    try:
        try:
            resp = opener.open(req, timeout=30)
        except urllib.error.HTTPError as exc:
            if exc.code in (301, 302, 303, 307, 308) and exc.headers.get("Location"):
                resp = urllib.request.urlopen(exc.headers["Location"], timeout=30)
            else:
                return None, f"Error HTTP {exc.code} al descargar el archivo."
        with resp:
            data = resp.read(_MAX_DOWNLOAD_BYTES + 1)
    except Exception as exc:  # noqa: BLE001  # HTTP request catch-all
        return None, f"No se pudo descargar el archivo: {exc}"
    if len(data) > _MAX_DOWNLOAD_BYTES:
        return None, f"El archivo supera el límite de {_MAX_DOWNLOAD_BYTES // (1024 * 1024)} MB."
    return data, ""


def _extract_docx(data: bytes, limit: int) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    total = 0
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
            total += len(para.text)
        if total > limit:
            return "\n".join(parts)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            total += len(parts[-1])
            if total > limit:
                return "\n".join(parts)
    return "\n".join(parts)


def _extract_xlsx(data: bytes, limit: int) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    total = 0
    for ws in wb.worksheets:
        parts.append(f"## Hoja: {ws.title}")
        for index, row in enumerate(ws.iter_rows(values_only=True)):
            if index >= _MAX_SHEET_ROWS or total > limit:
                break
            line = " | ".join("" if v is None else str(v) for v in row).rstrip(" |")
            if line:
                parts.append(line)
                total += len(line)
        if total > limit:
            break
    return "\n".join(parts)


def _extract_pptx(data: bytes, limit: int) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    total = 0
    for number, slide in enumerate(prs.slides, start=1):
        parts.append(f"## Diapositiva {number}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
                total += len(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
                    total += len(parts[-1])
        if total > limit:
            break
    return "\n".join(parts)


def _extract_pdf(data: bytes, limit: int) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    total_pages = len(reader.pages)
    parts: list[str] = []
    total = 0
    for number, page in enumerate(reader.pages[:_MAX_PDF_PAGES], start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"--- Página {number} ---\n{text}")
            total += len(text)
        if total > limit:
            break
    return "\n".join(parts), total_pages


_LEGACY_FORMATS = {
    ".doc": ".docx",
    ".rtf": ".docx",
    ".odt": ".docx",
    ".xls": ".xlsx",
    ".xlsb": ".xlsx",
    ".ods": ".xlsx",
    ".ppt": ".pptx",
    ".odp": ".pptx",
}


def _convert_legacy(name: str, data: bytes) -> tuple[str, bytes] | str:
    ext = os.path.splitext(name.lower())[1]
    target = _LEGACY_FORMATS[ext]
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        return f"No hay conversor instalado para leer archivos {ext}."
    with tempfile.TemporaryDirectory() as folder:
        source = os.path.join(folder, "origen" + ext)
        with open(source, "wb") as handle:
            handle.write(data)
        try:
            subprocess.run(
                [office, "--headless", "--convert-to", target.lstrip("."), "--outdir", folder, source],
                capture_output=True,
                timeout=120,
                env=dict(os.environ, HOME=folder),
                check=False,
            )
        except subprocess.TimeoutExpired:
            return f"La conversión del archivo {ext} tardó más de 2 minutos."
        converted = os.path.join(folder, "origen" + target)
        if not os.path.isfile(converted):
            return f"No se pudo convertir el archivo {ext} para leerlo (¿está dañado o protegido con contraseña?)."
        with open(converted, "rb") as handle:
            return os.path.splitext(name)[0] + target, handle.read()


def _extract_text(name: str, data: bytes, limit: int) -> tuple[str | None, str, dict]:
    ext = os.path.splitext(name.lower())[1]
    extra: dict = {}
    try:
        if ext in _TEXT_EXTENSIONS:
            return data.decode("utf-8", errors="replace"), "", extra
        if ext == ".docx":
            return _extract_docx(data, limit), "", extra
        if ext in (".xlsx", ".xlsm"):
            return _extract_xlsx(data, limit), "", extra
        if ext == ".pptx":
            return _extract_pptx(data, limit), "", extra
        if ext == ".pdf":
            text, pages = _extract_pdf(data, limit)
            extra["pages"] = pages
            if not text.strip():
                return None, "El PDF no contiene texto extraíble (puede ser un escaneo o imágenes).", extra
            return text, "", extra
    except ImportError as exc:
        return None, f"Falta una librería para leer archivos {ext}: {exc}", extra
    except Exception as exc:  # noqa: BLE001  # parser catch-all
        return None, f"No se pudo leer el contenido del archivo {ext}: {exc}", extra
    if ext in _LEGACY_FORMATS:
        converted = _convert_legacy(name, data)
        if isinstance(converted, str):
            return None, converted, extra
        text, error, extra = _extract_text(converted[0], converted[1], limit)
        extra["convertido_desde"] = ext
        return text, error, extra
    return None, f"Formato {ext or 'desconocido'} no soportado para lectura de contenido.", extra


def _save_to_workspace(name: str, data: bytes) -> str:
    safe = re.sub(r"[^\w.\- ]", "_", os.path.basename(name)).strip(" .") or "archivo"
    target_dir = Path(os.getcwd()) / _DOWNLOAD_DIR
    target_dir.mkdir(parents=True, exist_ok=True)
    dest = target_dir / f"{uuid.uuid4().hex[:8]}_{safe}"
    dest.write_bytes(data)
    return f"{_DOWNLOAD_DIR}/{dest.name}"


def _fetch_item(args: dict) -> tuple[str | None, dict]:
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."}), {}
    token, user_site_url, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err), {}

    file_path = str(args.get("file_path") or "").strip().strip("/")
    item_id = str(args.get("item_id") or "").strip()
    drive_id = str(args.get("drive_id") or "").strip()
    site_url = str(args.get("site_url") or user_site_url or "").strip().rstrip("/")

    if item_id:
        item_path = f"/drives/{drive_id}/items/{item_id}" if drive_id else f"/me/drive/items/{item_id}"
    elif file_path:
        if site_url:
            parsed = urllib.parse.urlparse(site_url)
            hostname = parsed.netloc
            site_path = parsed.path.rstrip("/") or "/"
            item_path = f"/sites/{hostname}:{site_path}:/drive/root:/{urllib.parse.quote(file_path)}"
        else:
            item_path = f"/me/drive/root:/{urllib.parse.quote(file_path)}"
    else:
        return json.dumps({"success": False, "error": "Se requiere file_path o item_id."}), {}

    meta = _graph("GET", item_path, token)
    if "error" in meta:
        return json.dumps({"success": False, "error": _with_hint(meta["error"])}, ensure_ascii=False), {}
    if "folder" in meta:
        return json.dumps({"success": False, "error": "La ruta es una carpeta; usa sharepoint_list_files."}), {}

    name = str(meta.get("name") or file_path or item_id)
    size = int(meta.get("size") or 0)
    if size > _MAX_DOWNLOAD_BYTES:
        message = (
            f"El archivo pesa {size / (1024 * 1024):.1f} MB; el límite es {_MAX_DOWNLOAD_BYTES // (1024 * 1024)} MB."
        )
        return json.dumps({"success": False, "error": message}), {}

    download_url = meta.get("@microsoft.graph.downloadUrl")
    if download_url:
        data, derr = _download_bytes(download_url)
    else:
        data, derr = _download_bytes(f"{GRAPH_BASE}{item_path.replace(' ', '%20')}/content", token)
    if data is None:
        return json.dumps({"success": False, "error": derr}), {}
    return None, {"name": name, "size": size, "url": meta.get("webUrl"), "data": data}


def _local_file_result(info: dict) -> dict:
    local_path = _save_to_workspace(info["name"], info["data"])
    ext = os.path.splitext(info["name"].lower())[1]
    result = {
        "success": True,
        "name": info["name"],
        "size": info["size"],
        "url": info["url"],
        "local_path": local_path,
        "mime_type": mimetypes.guess_type(info["name"])[0] or "application/octet-stream",
    }
    if ext in _IMAGE_EXTENSIONS:
        result["hint"] = "Imagen guardada en el workspace. Analízala con vision_analyze usando local_path, de a una."
    return result


_SHARE_FOLDER = "UltraBot-archivos"
_MAX_SHARE_BYTES = 50_000_000
_SIMPLE_UPLOAD_MAX_BYTES = 3_900_000
_UPLOAD_CHUNK_BYTES = 327_680 * 10
_SHARE_SCOPES = ("organization", "anonymous")


def _put_binary(url: str, data: bytes, token: str | None = None, content_range: str | None = None) -> dict:
    headers = {"Content-Type": "application/octet-stream", "Content-Length": str(len(data))}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    if content_range:
        headers["Content-Range"] = content_range
    req = urllib.request.Request(url, data=data, method="PUT", headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
            try:
                parsed = json.loads(raw) if raw else {}
            except ValueError:
                parsed = {}
            return parsed if isinstance(parsed, dict) else {}
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:300]
        if exc.code == 403:
            return {
                "error": "Sin permiso para escribir en OneDrive (403): la cuenta necesita el permiso Files.ReadWrite."
            }
        return {"error": f"Error {exc.code} subiendo el archivo: {detail}"}
    except (urllib.error.URLError, TimeoutError) as exc:
        return {"error": f"No se pudo subir el archivo: {exc}"}


def _resolve_local_file(raw_path: str) -> tuple[Path | None, str | None]:
    base = Path(os.environ.get("HERMES_SESSION_CWD") or os.getcwd()).resolve()
    candidate = Path(raw_path)
    path = (candidate if candidate.is_absolute() else base / candidate).resolve()
    if base != path and base not in path.parents:
        return None, f"El archivo '{raw_path}' está fuera del directorio de trabajo."
    if not path.is_file():
        return None, f"No existe el archivo '{raw_path}'."
    size = path.stat().st_size
    if size > _MAX_SHARE_BYTES:
        return None, f"El archivo pesa {size / 1_000_000:.1f} MB; el máximo es {_MAX_SHARE_BYTES // 1_000_000} MB."
    return path, None


def _safe_drive_name(display: str) -> str:
    return re.sub(r"[^A-Za-z0-9._ -]", "_", display).strip(" .")[:120] or "archivo"


def _upload_local_file(token: str, path: Path, remote_dir: str, safe_name: str) -> tuple[dict | None, str | None]:
    size = path.stat().st_size
    data = path.read_bytes()
    remote = f"/me/drive/root:/{remote_dir}/{safe_name}" if remote_dir else f"/me/drive/root:/{safe_name}"
    if size <= _SIMPLE_UPLOAD_MAX_BYTES:
        item = _put_binary(
            _encode_url(f"{GRAPH_BASE}{remote}:/content?@microsoft.graph.conflictBehavior=rename"), data, token
        )
    else:
        session = _graph(
            "POST", f"{remote}:/createUploadSession", token, {"item": {"@microsoft.graph.conflictBehavior": "rename"}}
        )
        if "error" in session:
            return None, session["error"]
        upload_url = session.get("uploadUrl")
        if not upload_url:
            return None, "OneDrive no entregó la dirección de subida."
        item = {}
        for start in range(0, size, _UPLOAD_CHUNK_BYTES):
            chunk = data[start : start + _UPLOAD_CHUNK_BYTES]
            item = _put_binary(upload_url, chunk, None, f"bytes {start}-{start + len(chunk) - 1}/{size}")
            if "error" in item:
                break
    if "error" in item:
        return None, item["error"]
    return item, None


_WINDOWS_ONEDRIVE_RE = re.compile(r"^[A-Za-z]:[\\/]+Users[\\/]+[^\\/]+[\\/]+OneDrive[^\\/]*[\\/]*(.*)$", re.I)


def _normalize_drive_folder(raw: str) -> tuple[str | None, str | None]:
    """Convierte una ruta de Windows o de OneDrive en una ruta relativa segura a la raíz del OneDrive."""
    text = str(raw or "").strip()
    match = _WINDOWS_ONEDRIVE_RE.match(text)
    if match:
        text = match.group(1)
    parts = [p.strip() for p in re.split(r"[\\/]+", text) if p.strip()]
    if len(parts) > 10:
        return None, "La carpeta tiene demasiados niveles (máximo 10)."
    for part in parts:
        if part in (".", "..") or re.search(r'[<>:"|?*\x00-\x1f]', part):
            return None, f"Nombre de carpeta no válido: '{part}'."
    return "/".join(parts), None


def _upload_file_tool(args: dict, **_kwargs) -> str:
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."})
    token, _site, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err)
    raw_path = str(args.get("file_path") or "").strip()
    if not raw_path:
        return json.dumps({"success": False, "error": "Se requiere file_path (archivo del directorio de trabajo)."})
    folder, folder_error = _normalize_drive_folder(args.get("folder"))
    if folder_error:
        return json.dumps({"success": False, "error": folder_error}, ensure_ascii=False)
    path, path_error = _resolve_local_file(raw_path)
    if path_error:
        return json.dumps({"success": False, "error": path_error}, ensure_ascii=False)
    display = str(args.get("name") or path.name).strip() or path.name
    item, upload_error = _upload_local_file(token, path, folder, _safe_drive_name(display))
    if upload_error:
        return json.dumps({"success": False, "error": upload_error}, ensure_ascii=False)
    return json.dumps(
        {
            "success": True,
            "name": item.get("name"),
            "folder": folder or "(raíz de OneDrive)",
            "size": path.stat().st_size,
            "url": item.get("webUrl"),
            "note": "Si OneDrive está sincronizado en su equipo, el archivo aparece solo en esa carpeta. Si ya existía uno con el mismo nombre, se guardó con otro nombre; nada se sobrescribe.",
        },
        ensure_ascii=False,
    )


def _share_file_tool(args: dict, **_kwargs) -> str:
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."})
    token, _site, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err)
    raw_path = str(args.get("file_path") or "").strip()
    if not raw_path:
        return json.dumps({"success": False, "error": "Se requiere file_path (archivo del directorio de trabajo)."})
    scope = str(args.get("scope") or "organization").strip().lower()
    if scope not in _SHARE_SCOPES:
        return json.dumps({"success": False, "error": "scope debe ser organization o anonymous."})
    base = Path(os.environ.get("HERMES_SESSION_CWD") or os.getcwd()).resolve()
    candidate = Path(raw_path)
    path = (candidate if candidate.is_absolute() else base / candidate).resolve()
    if base != path and base not in path.parents:
        return json.dumps({"success": False, "error": f"El archivo '{raw_path}' está fuera del directorio de trabajo."})
    if not path.is_file():
        return json.dumps({"success": False, "error": f"No existe el archivo '{raw_path}'."})
    size = path.stat().st_size
    if size > _MAX_SHARE_BYTES:
        return json.dumps(
            {
                "success": False,
                "error": f"El archivo pesa {size / 1_000_000:.1f} MB; el máximo es {_MAX_SHARE_BYTES // 1_000_000} MB.",
            }
        )
    display = str(args.get("name") or path.name).strip() or path.name
    safe = re.sub(r"[^A-Za-z0-9._ -]", "_", display).strip(" .")[:120] or "archivo"
    data = path.read_bytes()
    remote = f"/me/drive/root:/{_SHARE_FOLDER}/{safe}"
    if size <= _SIMPLE_UPLOAD_MAX_BYTES:
        item = _put_binary(
            _encode_url(f"{GRAPH_BASE}{remote}:/content?@microsoft.graph.conflictBehavior=rename"), data, token
        )
    else:
        session = _graph(
            "POST", f"{remote}:/createUploadSession", token, {"item": {"@microsoft.graph.conflictBehavior": "rename"}}
        )
        if "error" in session:
            return json.dumps({"success": False, "error": session["error"]}, ensure_ascii=False)
        upload_url = session.get("uploadUrl")
        if not upload_url:
            return json.dumps({"success": False, "error": "OneDrive no entregó la dirección de subida."})
        item = {}
        for start in range(0, size, _UPLOAD_CHUNK_BYTES):
            chunk = data[start : start + _UPLOAD_CHUNK_BYTES]
            item = _put_binary(upload_url, chunk, None, f"bytes {start}-{start + len(chunk) - 1}/{size}")
            if "error" in item:
                break
    if "error" in item:
        return json.dumps({"success": False, "error": item["error"]}, ensure_ascii=False)
    item_id = item.get("id")
    if not item_id:
        return json.dumps({"success": False, "error": "OneDrive no devolvió el identificador del archivo."})
    link = _graph("POST", f"/me/drive/items/{item_id}/createLink", token, {"type": "view", "scope": scope})
    if "error" in link:
        return json.dumps({"success": False, "error": link["error"]}, ensure_ascii=False)
    url = (link.get("link") or {}).get("webUrl")
    if not url:
        return json.dumps({"success": False, "error": "OneDrive no devolvió el enlace para compartir."})
    result = {"success": True, "name": item.get("name") or safe, "size": size, "url": url, "scope": scope}
    if scope == "organization":
        result["note"] = "Enlace solo para personas de la organización; destinatarios externos no podrán abrirlo."
    return json.dumps(result, ensure_ascii=False)


def _download_file_tool(args: dict, **_kwargs) -> str:
    error, info = _fetch_item(args)
    if error:
        return error
    return json.dumps(_local_file_result(info), ensure_ascii=False)


_FOLDER_DEFAULT_FILES = 40


_FOLDER_HARD_FILES = 100


_FOLDER_DEFAULT_MB = 150


_FOLDER_HARD_MB = 300


_FOLDER_MAX_DEPTH = 3


_FOLDER_WORKERS = 4


def _clean_extensions(raw) -> set[str]:
    values = raw if isinstance(raw, list) else ([raw] if raw else [])
    return {"." + str(v).strip().lower().lstrip(".") for v in values if str(v).strip()}


def _collect_folder_files(
    base: str, folder: str, token: str, recursive: bool, depth: int, relative: str, found: list, errors: list
) -> None:
    items, error = _fetch_children(base, folder, token)
    if items is None:
        errors.append(f"{relative or '/'}: {str(error)[:80]}")
        return
    for item in items:
        name = str(item.get("name") or "")
        if "folder" in item:
            if recursive and depth < _FOLDER_MAX_DEPTH and name:
                child_relative = f"{relative}/{name}" if relative else name
                parent_drive = (item.get("parentReference") or {}).get("driveId")
                if folder == "" and base.startswith("/drives/") and item.get("id") and parent_drive:
                    child_base, child_folder = f"/drives/{parent_drive}/items/{item['id']}", ""
                else:
                    child_base, child_folder = base, f"{folder}/{name}" if folder else name
                _collect_folder_files(
                    child_base, child_folder, token, recursive, depth + 1, child_relative, found, errors
                )
        else:
            found.append({"item": item, "relative": relative})


def _download_folder_tool(args: dict, **_kwargs) -> str:
    user_id = _task_user_id()
    if not user_id:
        return json.dumps({"success": False, "error": "No se pudo determinar el usuario de esta tarea."})
    token, user_site_url, err = _get_m365_token(user_id)
    if not token:
        return _auth_error(err)

    site_url = str(args.get("site_url") or user_site_url or "").strip().rstrip("/")
    folder_path = str(args.get("folder_path") or "").strip().strip("/")
    item_id = str(args.get("item_id") or "").strip()
    drive_id = str(args.get("drive_id") or "").strip()
    extensions = _clean_extensions(args.get("extensions"))
    recursive = bool(args.get("recursive"))
    try:
        max_files = max(1, min(_FOLDER_HARD_FILES, int(args.get("max_files") or _FOLDER_DEFAULT_FILES)))
        max_bytes = max(1, min(_FOLDER_HARD_MB, int(args.get("max_mb") or _FOLDER_DEFAULT_MB))) * 1024 * 1024
    except (TypeError, ValueError):
        max_files, max_bytes = _FOLDER_DEFAULT_FILES, _FOLDER_DEFAULT_MB * 1024 * 1024

    if item_id:
        label = item_id
        item_base = f"/drives/{drive_id}/items/{item_id}" if drive_id else f"/me/drive/items/{item_id}"
        candidates = [(item_base, "", None)]
    elif folder_path:
        label = folder_path
        if not site_url:
            folder_path, folder_error = _normalize_drive_folder(folder_path)
            if folder_error:
                return json.dumps({"success": False, "error": folder_error}, ensure_ascii=False)
        candidates = _site_bases(site_url, folder_path, token) if site_url else [("/me/drive/root", folder_path, None)]
    else:
        return json.dumps({"success": False, "error": "Indica folder_path (o item_id de la carpeta)."})

    found: list = []
    errors: list = []
    last_error = "No se pudo leer la carpeta."
    for base, folder, _name in candidates:
        found, errors = [], []
        probe, probe_error = _fetch_children(base, folder, token)
        if probe is None:
            last_error = probe_error or last_error
            continue
        _collect_folder_files(base, folder, token, recursive, 1, "", found, errors)
        break
    else:
        return json.dumps({"success": False, "error": _with_hint(last_error)}, ensure_ascii=False)

    selected: list = []
    skipped: list = []
    total = 0
    for entry in found:
        item = entry["item"]
        name = str(item.get("name") or "archivo")
        size = int(item.get("size") or 0)
        ext = os.path.splitext(name.lower())[1]
        if extensions and ext not in extensions:
            continue
        if size > _MAX_DOWNLOAD_BYTES:
            skipped.append({"name": name, "motivo": f"pesa {size / (1024 * 1024):.1f} MB (máximo 10 MB por archivo)"})
        elif len(selected) >= max_files:
            skipped.append({"name": name, "motivo": f"se alcanzó el máximo de {max_files} archivos"})
        elif total + size > max_bytes:
            skipped.append(
                {"name": name, "motivo": f"se alcanzó el máximo de {max_bytes // (1024 * 1024)} MB en total"}
            )
        else:
            total += size
            selected.append(entry)

    slug = re.sub(r"[^\w.\- ]", "_", os.path.basename(str(label).rstrip("/"))).strip(" .") or "carpeta"
    destination = Path(os.getcwd()) / _DOWNLOAD_DIR / slug
    destination.mkdir(parents=True, exist_ok=True)

    def _fetch_one(entry: dict) -> dict:
        item = entry["item"]
        name = str(item.get("name") or "archivo")
        url = item.get("@microsoft.graph.downloadUrl")
        if url:
            data, derr = _download_bytes(url)
        else:
            drive = (item.get("parentReference") or {}).get("driveId") or drive_id
            path = (
                f"/drives/{drive}/items/{item.get('id')}/content"
                if drive
                else f"/me/drive/items/{item.get('id')}/content"
            )
            data, derr = _download_bytes(f"{GRAPH_BASE}{path}", token)
        if data is None:
            return {"name": name, "error": derr}
        relative = entry["relative"]
        safe = re.sub(r"[^\w.\- ]", "_", name).strip(" .") or "archivo"
        prefix = re.sub(r"[^\w.\- ]", "_", relative.replace("/", "__")) + "__" if relative else ""
        target = destination / f"{prefix}{safe}"
        target.write_bytes(data)
        return {
            "name": name,
            "size": len(data),
            "local_path": f"{_DOWNLOAD_DIR}/{slug}/{target.name}",
            "subcarpeta": relative or None,
        }

    downloaded: list = []
    with ThreadPoolExecutor(max_workers=_FOLDER_WORKERS) as pool:
        for result in pool.map(_fetch_one, selected):
            if "error" in result:
                skipped.append({"name": result["name"], "motivo": result["error"]})
            else:
                downloaded.append(result)

    output = {
        "success": True,
        "carpeta": label,
        "descargados": len(downloaded),
        "archivos": downloaded,
        "omitidos": skipped,
        "bytes_totales": sum(f["size"] for f in downloaded),
        "hint": "Archivos en el workspace. Lee los de texto con read_file y analiza las imágenes con vision_analyze, de a una.",
    }
    if errors:
        output["subcarpetas_sin_leer"] = errors[:5]
    return json.dumps(output, ensure_ascii=False)


def _read_file_tool(args: dict, **_kwargs) -> str:
    try:
        max_chars = int(args.get("max_chars") or _DEFAULT_MAX_CHARS)
    except (TypeError, ValueError):
        max_chars = _DEFAULT_MAX_CHARS
    max_chars = max(500, min(max_chars, _HARD_MAX_CHARS))

    error, info = _fetch_item(args)
    if error:
        return error
    name, size, data = info["name"], info["size"], info["data"]

    if os.path.splitext(name.lower())[1] in _IMAGE_EXTENSIONS:
        return json.dumps(_local_file_result(info), ensure_ascii=False)

    text, terr, extra = _extract_text(name, data, max_chars)
    if text is None:
        return json.dumps({"success": False, "name": name, "size": size, "error": terr, **extra})

    truncated = len(text) > max_chars
    return json.dumps(
        {
            "success": True,
            "name": name,
            "size": size,
            "url": info["url"],
            "truncated": truncated,
            "chars": min(len(text), max_chars),
            "content": text[:max_chars],
            **extra,
        },
        ensure_ascii=False,
    )


def register(ctx):
    ctx.register_tool(
        name="sharepoint_list_drives",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_list_drives",
            "description": "Lista las unidades de almacenamiento accesibles: OneDrive personal y bibliotecas de documentos de SharePoint.",
            "parameters": {"type": "object", "properties": {}},
        },
        handler=_list_drives_tool,
        description="Listar drives disponibles (OneDrive + SharePoint)",
        emoji="🗄️",
    )
    ctx.register_tool(
        name="sharepoint_list_files",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_list_files",
            "description": (
                "Lista archivos y carpetas en OneDrive o en un sitio SharePoint. Con site_url y folder_path prueba solo las "
                "bibliotecas del sitio hasta encontrar la carpeta (no hace falta adivinar el nombre de la biblioteca). "
                "Con depth=2 o 3 devuelve el ÁRBOL de subcarpetas en una sola llamada: úsalo para mapear o entender cómo "
                "está estructurada una carpeta, en vez de buscar con sharepoint_search una y otra vez. Para «mostrar todas "
                "las carpetas a las que tengo acceso» llama UNA vez con solo_carpetas=true y depth=4 (devuelve solo carpetas, "
                "con cuántos archivos tiene cada una); NO llames carpeta por carpeta."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "site_url": {
                        "type": "string",
                        "description": "URL del sitio SharePoint (opcional, ej: https://empresa.sharepoint.com/sites/Marketing). Si no se indica, usa el OneDrive del usuario.",
                    },
                    "folder_path": {
                        "type": "string",
                        "description": "Ruta de carpeta (opcional, ej: 'Gerencia/TI/10. SGSI' o 'Documentos/Proyectos')",
                    },
                    "depth": {
                        "type": "integer",
                        "description": "Niveles de subcarpetas a devolver como árbol (1 a 3, default 1; hasta 4 con solo_carpetas)",
                    },
                    "solo_carpetas": {
                        "type": "boolean",
                        "description": "true: el árbol trae solo carpetas (con el número de archivos de cada una), mucho más corto; ideal para mapear todo.",
                    },
                    "max_items": {
                        "type": "integer",
                        "description": "Máximo de elementos del árbol (50 a 800, default 300).",
                    },
                },
            },
        },
        handler=_list_files_tool,
        description="Listar archivos en SharePoint/OneDrive",
        emoji="📁",
    )
    ctx.register_tool(
        name="sharepoint_get_file",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_get_file",
            "description": "Obtiene información de un archivo por su ruta en OneDrive o SharePoint.",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {"type": "string", "description": "Ruta del archivo (ej: 'Documents/informe.pdf')"},
                    "site_url": {"type": "string", "description": "URL del sitio SharePoint (opcional)"},
                },
                "required": ["file_path"],
            },
        },
        handler=_get_file_tool,
        description="Obtener archivo de SharePoint/OneDrive",
        emoji="📄",
    )
    ctx.register_tool(
        name="sharepoint_search",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_search",
            "description": (
                "Busca archivos y documentos en SharePoint y OneDrive usando Microsoft Search. Cada resultado trae "
                "item_id y drive_id: para leer o descargar un resultado pásalos a sharepoint_read_file o "
                "sharepoint_download_file en vez de adivinar la ruta o el site_url."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string", "description": "Término de búsqueda"},
                    "count": {"type": "integer", "description": "Número de resultados (máx 25, default 10)"},
                },
                "required": ["query"],
            },
        },
        handler=_search_tool,
        description="Buscar archivos en SharePoint/OneDrive",
        emoji="🔍",
    )
    ctx.register_tool(
        name="sharepoint_read_file",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_read_file",
            "description": (
                "Lee el CONTENIDO de texto de un archivo de OneDrive o SharePoint "
                "(PDF, Word .docx, Excel .xlsx, PowerPoint .pptx, txt, csv, md, json, html). "
                "Si el archivo es una imagen, la guarda en el workspace y devuelve local_path para analizarla "
                "con vision_analyze. Usa esta herramienta cuando el usuario pida saber qué dice o contiene un "
                "archivo. Máximo 10 MB; el texto se recorta a max_chars."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "Ruta del archivo (ej: 'Documentos/informe.pdf'). Alternativa a item_id.",
                    },
                    "item_id": {
                        "type": "string",
                        "description": "ID del archivo (de sharepoint_search o sharepoint_list_files). Alternativa a file_path.",
                    },
                    "drive_id": {"type": "string", "description": "ID de la unidad (opcional, junto con item_id)"},
                    "site_url": {"type": "string", "description": "URL del sitio SharePoint (opcional)"},
                    "max_chars": {
                        "type": "integer",
                        "description": "Máximo de caracteres a devolver (default 20000, tope 100000)",
                    },
                },
            },
        },
        handler=_read_file_tool,
        description="Leer contenido de un archivo de SharePoint/OneDrive",
        emoji="📖",
    )
    ctx.register_tool(
        name="sharepoint_share_file",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_share_file",
            "description": (
                "Sube un archivo del directorio de trabajo del agente (PDF, PPTX, imagen, etc.; máx. 50 MB) al OneDrive "
                "del usuario (carpeta UltraBot-archivos) y devuelve un ENLACE para abrirlo o compartirlo. Úsala cuando "
                "el usuario pida el link de un archivo generado o cuando no pueda visualizarlo en el chat."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "Ruta del archivo dentro del directorio de trabajo (ej. outputs/informe.pdf)",
                    },
                    "name": {"type": "string", "description": "Nombre con que se guarda (opcional)"},
                    "scope": {
                        "type": "string",
                        "enum": ["organization", "anonymous"],
                        "description": (
                            "organization (default): solo personas de la organización. anonymous: cualquiera con el "
                            "enlace; úsalo SOLO si el usuario pide compartir con externos y su organización lo permite."
                        ),
                    },
                },
                "required": ["file_path"],
            },
        },
        handler=_share_file_tool,
        description="Subir archivo a OneDrive y obtener enlace",
        emoji="🔗",
    )
    ctx.register_tool(
        name="sharepoint_upload_file",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_upload_file",
            "description": (
                "Guarda un archivo del directorio de trabajo del agente (máx. 50 MB) en una carpeta del OneDrive "
                "del propio usuario, por ejemplo Escritorio/Proyectos/IA. Acepta una ruta de Windows como "
                "C:\\Users\\x\\OneDrive - Contoso\\Escritorio\\Proyectos (se traduce sola). Crea las carpetas "
                "que falten y NUNCA sobrescribe: si ya existe un archivo con ese nombre, lo guarda con otro. "
                "Si el OneDrive está sincronizado, el archivo aparece en el equipo del usuario. Úsala cuando el "
                "usuario pida guardar o dejar un archivo en una carpeta suya; para solo un enlace usa sharepoint_share_file."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "Ruta del archivo dentro del directorio de trabajo (ej. outputs/informe.docx)",
                    },
                    "folder": {
                        "type": "string",
                        "description": "Carpeta de destino en el OneDrive del usuario (ej. Escritorio/Usuario/Proyectos/IA) o su ruta de Windows",
                    },
                    "name": {"type": "string", "description": "Nombre con que se guarda (opcional)"},
                },
                "required": ["file_path", "folder"],
            },
        },
        handler=_upload_file_tool,
        description="Guardar archivo en una carpeta del OneDrive del usuario",
        emoji="📤",
    )
    ctx.register_tool(
        name="sharepoint_download_file",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_download_file",
            "description": (
                "Descarga un archivo de OneDrive o SharePoint al workspace del agente y devuelve su local_path. "
                "Úsala para imágenes (luego vision_analyze con local_path) o para procesar el archivo con código. "
                "No uses la URL web de SharePoint con otras herramientas: pide inicio de sesión. Máximo 10 MB."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "Ruta del archivo (ej: 'Documentos/foto.png'). Alternativa a item_id.",
                    },
                    "item_id": {"type": "string", "description": "ID del archivo. Alternativa a file_path."},
                    "drive_id": {"type": "string", "description": "ID de la unidad (opcional, junto con item_id)"},
                    "site_url": {"type": "string", "description": "URL del sitio SharePoint (opcional)"},
                },
            },
        },
        handler=_download_file_tool,
        description="Descargar archivo de SharePoint/OneDrive al workspace",
        emoji="⬇️",
    )
    ctx.register_tool(
        name="sharepoint_download_folder",
        toolset=TOOLSET,
        schema={
            "name": "sharepoint_download_folder",
            "description": (
                "Descarga TODOS los archivos de una carpeta de OneDrive o SharePoint al workspace en UNA llamada "
                "(en paralelo) y devuelve la lista con su local_path. Úsala cuando necesites varios archivos de una "
                "carpeta (una biblioteca de referencia, un lote de imágenes o documentos): NO llames a "
                "sharepoint_search y sharepoint_download_file archivo por archivo. Filtra por extensiones y, con "
                "recursive=true, entra en subcarpetas (hasta 3 niveles). Máximo 100 archivos y 300 MB por llamada; "
                "cada archivo hasta 10 MB. Los que no caben salen en 'omitidos' con su motivo."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "folder_path": {
                        "type": "string",
                        "description": "Ruta de la carpeta (ej: 'Escritorio/Proyectos/IA/Referencias'). Alternativa a item_id.",
                    },
                    "site_url": {"type": "string", "description": "URL del sitio SharePoint (opcional)"},
                    "item_id": {
                        "type": "string",
                        "description": "ID de la carpeta (de sharepoint_search). Alternativa a folder_path.",
                    },
                    "drive_id": {"type": "string", "description": "ID de la unidad (opcional, junto con item_id)"},
                    "extensions": {
                        "type": "array",
                        "items": {"type": "string"},
                        "description": "Solo estas extensiones, ej. ['pptx','png']. Sin indicar, baja todo.",
                    },
                    "recursive": {"type": "boolean", "description": "true: incluye subcarpetas (hasta 3 niveles)."},
                    "max_files": {"type": "integer", "description": "Máximo de archivos (default 40, tope 100)."},
                    "max_mb": {"type": "integer", "description": "Máximo de MB en total (default 150, tope 300)."},
                },
            },
        },
        handler=_download_folder_tool,
        description="Descargar una carpeta completa de SharePoint/OneDrive al workspace",
        emoji="📥",
    )
